"""Exporta o orçamento como planilha Excel EDITÁVEL — ao contrário do
relatório de impressão (que é só leitura), aqui os valores importantes
(peso, taxas por processo, alíquotas) ficam em células próprias (painel
"Parâmetros" à direita, colunas I:J) e o relatório em si (colunas A:G)
referencia essas células por fórmula, não por valor fixo. Mudar um
parâmetro recalcula tudo — é a mesma matemática de app/processos.py e
app/comercial.py, só que reescrita como fórmula do Excel em vez de Python.

Tudo numa aba só (pedido do usuário): o relatório (A:G) é o que aparece
ao imprimir/exportar em PDF do Excel (`print_area` cobre só essas
colunas, layout A4 retrato); o painel de parâmetros fica ao lado, visível
e editável na mesma aba, mas fora da área de impressão.

O que NÃO vira fórmula (fica como número editável direto, sem derivação):
usinagem (lista de operações ad hoc) e itens_padrao (idem) — não têm uma
taxa única por kg/peça pra expor como parâmetro, e horas de caldeiraria
quando vêm do histórico combinado (usar_historico_horas) em vez da regra
simples peso×fator.
"""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagemExcel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.page import PageMargins

from app.parametros_padrao import ALIQUOTAS_COMPRA_POR_TIPO, PARAMETROS_PADRAO

# Empresa que emite o orçamento (pedido explícito do usuário) — logo
# extraída da ficha cadastral real da Macfab. "FC Nexus" é só o nome do
# software, por isso vira uma marca pequena e discreta no cabeçalho, não
# o título principal da planilha.
_MACFAB_NOME = "MACFAB Fabricações e Serviços Industriais LTDA"
_MACFAB_CNPJ = "13.014.242/0001-86"
_CAMINHO_LOGO_MACFAB = Path(__file__).resolve().parent.parent / "data" / "assets" / "macfab-logo.png"

MOEDA = '"R$" #,##0.00'
PERCENTUAL = "0.00%"

_AZUL_ESCURO = "0F172A"
_AZUL_MEDIO = "1E293B"
_CIANO = "0891B2"
_CINZA_CLARO = "E2E8F0"
_BRANCO = "FFFFFF"

_FONTE_TITULO = Font(color=_BRANCO, bold=True, size=16)
_FONTE_SUBTITULO = Font(color=_CINZA_CLARO, size=10, italic=True)
_FONTE_MARCA_DISCRETA = Font(color=_CINZA_CLARO, size=7, italic=True)
_FONTE_SECAO = Font(color=_BRANCO, bold=True, size=11)
_FONTE_CABECALHO_TABELA = Font(bold=True, size=10)
_FONTE_TOTAL = Font(bold=True, size=10)
_FONTE_RESUMO_ROTULO = Font(size=10)
_FONTE_RESUMO_VALOR = Font(bold=True, size=10)
_FONTE_PRECO_FINAL = Font(color=_BRANCO, bold=True, size=15)
_FONTE_PARAMETROS_TITULO = Font(bold=True, size=10, color=_BRANCO)

_FILL_TITULO = PatternFill(start_color=_AZUL_ESCURO, end_color=_AZUL_ESCURO, fill_type="solid")
_FILL_SECAO = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")
_FILL_CABECALHO_TABELA = PatternFill(start_color=_CINZA_CLARO, end_color=_CINZA_CLARO, fill_type="solid")
_FILL_PRECO_FINAL = PatternFill(start_color=_CIANO, end_color=_CIANO, fill_type="solid")
_FILL_PARAMETROS_TITULO = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")

_BORDA_FINA = Border(*(Side(style="thin", color="CBD5E1") for _ in range(4)))

_COL_PARAM_VALOR = "J"


def _mesclar_e_estilizar(ws, linha: int, col_ini: str, col_fim: str, valor, fonte: Font, fill: PatternFill, alinhamento: str = "left") -> None:
    ws.merge_cells(f"{col_ini}{linha}:{col_fim}{linha}")
    celula = ws[f"{col_ini}{linha}"]
    celula.value = valor
    celula.font = fonte
    celula.fill = fill
    celula.alignment = Alignment(horizontal=alinhamento, vertical="center")
    for indice_col in range(column_index_from_string(col_ini), column_index_from_string(col_fim) + 1):
        ws[f"{get_column_letter(indice_col)}{linha}"].fill = fill


def _cabecalho_tabela(ws, linha: int, titulos: list[str]) -> None:
    for i, titulo in enumerate(titulos):
        celula = ws.cell(row=linha, column=i + 1, value=titulo)
        celula.font = _FONTE_CABECALHO_TABELA
        celula.fill = _FILL_CABECALHO_TABELA
        celula.border = _BORDA_FINA
        celula.alignment = Alignment(horizontal="left" if i == 0 else "center")


def _bordar_linha(ws, linha: int, n_colunas: int) -> None:
    for col in range(1, n_colunas + 1):
        ws.cell(row=linha, column=col).border = _BORDA_FINA


def _montar_parametros(ws, entrada: dict, params: dict, comercial) -> dict[str, int]:
    """Escreve o painel de parâmetros (colunas I:J) e devolve
    {chave: numero_da_linha} pro relatório referenciar por fórmula."""
    _mesclar_e_estilizar(
        ws, 1, "I", "J", "PARÂMETROS (editável)", _FONTE_PARAMETROS_TITULO, _FILL_PARAMETROS_TITULO, "center"
    )
    linhas: dict[str, int] = {}
    r = 2

    def add(chave: str, rotulo: str, valor, fmt: str | None = None):
        nonlocal r
        ws.cell(row=r, column=9, value=rotulo).font = Font(size=9)
        celula = ws.cell(row=r, column=10, value=valor)
        celula.font = Font(size=9, bold=True)
        if fmt:
            celula.number_format = fmt
        linhas[chave] = r
        r += 1

    add("peso_liquido_kg", "Peso líquido (kg) — usado no cálculo", entrada["peso_liquido_kg"], "0.00")
    # Só informativo (nenhuma fórmula referencia essa célula) — mostra o
    # peso calculado pelo sistema mesmo quando "Peso líquido (kg)" acima
    # foi trocado por um valor manual (ver PainelPesoBase.tsx no
    # frontend), pra quem abrir o Excel entender de onde veio a diferença.
    add(
        "peso_bruto_calculado_kg", "Peso bruto calculado (kg, referência)",
        entrada.get("peso_bruto_calculado_kg", entrada["peso_liquido_kg"]), "0.00",
    )
    add("fator_margem", "Fator de margem de venda", params["fator_margem_venda"], "0.00")
    add("corte_valor_kg", "Corte — R$/kg", params["corte_valor_kg"], MOEDA)
    add(
        "corte_peso_kg", "Corte — peso usado (kg, vazio = peso líquido)",
        params.get("corte_peso_kg") if params.get("corte_peso_kg") is not None else "",
        "0.00",
    )
    add(
        "corte_fator_percentual_adicional", "Corte — fator adicional (%)",
        params.get("corte_fator_percentual_adicional") or 0.0, "0.00",
    )
    add(
        "caldeiraria_peso_kg", "Caldeiraria — peso (kg, vazio = peso líquido)",
        params.get("caldeiraria_peso_kg") if params.get("caldeiraria_peso_kg") is not None else "",
        "0.00",
    )
    add("caldeiraria_fator_h_kg", "Caldeiraria — fator h/kg", params["caldeiraria_fator_h_kg"], "0.0000")
    add("caldeiraria_valor_hora", "Caldeiraria — R$/h", params["caldeiraria_valor_hora"], MOEDA)
    add("jateamento_divisor", "Jateamento/pintura — divisor sobre h caldeiraria", params["jateamento_pintura_divisor"], "0")
    add("jateamento_valor_hora", "Jateamento/pintura — R$/h", params["jateamento_pintura_valor_hora"], MOEDA)
    add(
        "solda_peso_kg", "Solda — peso do consumível (kg, vazio = peso líquido)",
        params.get("solda_peso_kg") if params.get("solda_peso_kg") is not None else "",
        "0.00",
    )
    add("solda_fator_consumo", "Solda — % consumível carbono sobre peso", params["solda_fator_consumo_percentual"], PERCENTUAL)
    add("solda_preco_consumivel", "Solda — R$/kg consumível carbono", params["solda_preco_kg_consumivel"], MOEDA)
    add(
        "solda_gas_peso_kg", "Solda — peso-base do gás (kg, vazio = kg de consumível)",
        params.get("solda_gas_peso_kg") if params.get("solda_gas_peso_kg") is not None else "",
        "0.00",
    )
    add("solda_fator_gas", "Solda — % gás (100% CO2) sobre consumível", params["solda_fator_gas_sobre_consumivel"], PERCENTUAL)
    add("solda_preco_gas", "Solda — R$/unidade gás (100% CO2)", params["solda_preco_unidade_gas"], MOEDA)
    add("solda_inox_qtd_kg", "Solda — kg consumível inox (0 = não usa)", params.get("solda_inox_qtd_kg") or 0.0, "0.00")
    add("solda_inox_preco_consumivel", "Solda — R$/kg consumível inox", params["solda_inox_preco_kg_consumivel"], MOEDA)
    add("solda_inox_fator_gas", "Solda — % gás (25% Ar-75% CO2) sobre consumível inox", params["solda_inox_fator_gas_sobre_consumivel"], PERCENTUAL)
    add("solda_inox_preco_gas", "Solda — R$/unidade gás (25% Ar-75% CO2)", params["solda_inox_preco_unidade_gas"], MOEDA)
    add("pintura_fator_l_m2", "Pintura — L/m² por demão", params["pintura_fator_l_m2"], "0.0000")
    add("area_pintura_m2", "Área de pintura (m²)", entrada.get("area_pintura_m2") or 0, "0.00")
    demaos = params["pintura_demaos"]
    preco_fundo = next((d["preco_litro"] for d in demaos if d["tipo"] == "fundo"), 0.0)
    preco_acabamento = next((d["preco_litro"] for d in demaos if d["tipo"] == "acabamento"), 0.0)
    add("pintura_preco_fundo", "Pintura — R$/L fundo", preco_fundo, MOEDA)
    add("pintura_preco_acabamento", "Pintura — R$/L acabamento", preco_acabamento, MOEDA)
    add("ndt_valor_kg", "NDT — R$/kg", params["ndt_valor_kg"], MOEDA)
    add("engenharia_valor_unitario", "Engenharia — R$/posição", params["engenharia_valor_unitario"], MOEDA)
    add("qtd_posicoes_engenharia", "Quantidade de posições de engenharia", entrada.get("quantidade_posicoes_engenharia") or 0, "0")
    add("embalagem_valor_kg", "Embalagem — R$/kg", params["embalagem_valor_kg"], MOEDA)
    add(
        "embalagem_peso_kg", "Embalagem — peso (kg, vazio = peso líquido)",
        params.get("embalagem_peso_kg") if params.get("embalagem_peso_kg") is not None else "",
        "0.00",
    )
    add("transporte_valor_kg", "Transporte — R$/kg", params["transporte_valor_kg"], MOEDA)
    add(
        "transporte_peso_kg", "Transporte — peso (kg, vazio = peso líquido)",
        params.get("transporte_peso_kg") if params.get("transporte_peso_kg") is not None else "",
        "0.00",
    )
    add("energia_valor_kg", "Energia — R$/kg", params["energia_valor_kg"], MOEDA)
    add(
        "energia_peso_kg", "Energia — peso (kg, vazio = peso líquido)",
        params.get("energia_peso_kg") if params.get("energia_peso_kg") is not None else "",
        "0.00",
    )
    add("aliquota_venda", f"Alíquota de venda ({comercial.cenario_comercial})", comercial.aliquota_venda, PERCENTUAL)

    ws.column_dimensions["I"].width = 40
    ws.column_dimensions["J"].width = 14
    return linhas


def _ref(linhas: dict[str, int], chave: str) -> str:
    return f"${_COL_PARAM_VALOR}${linhas[chave]}"


def _formula_bruto_processo(codigo: str, entrada: dict, linhas_p: dict[str, int], linha_caldeiraria_horas_cel: str | None) -> tuple[str, str | None]:
    """Devolve (formula_bruto, formula_horas_ou_None) pro código de processo."""
    peso = _ref(linhas_p, "peso_liquido_kg")

    if codigo == "corte":
        peso_corte_ref = _ref(linhas_p, "corte_peso_kg")
        peso_corte = f'IF({peso_corte_ref}="",{peso},{peso_corte_ref})'
        fator_ref = _ref(linhas_p, "corte_fator_percentual_adicional")
        return f"=({peso_corte})*{_ref(linhas_p, 'corte_valor_kg')}*(1+{fator_ref}/100)", None

    if codigo == "caldeiraria":
        if entrada.get("usar_historico_horas"):
            return None, None  # tratado como valor estático fora daqui
        # app/processos.py calcula o bruto com as horas SEM arredondar, mas
        # guarda linha.horas arredondado (2 casas) — e é esse valor
        # arredondado que orcamento.py repassa pro jateamento/pintura. Pra
        # bater exato com o motor real, o bruto usa a expressão cheia (sem
        # arredondar) e a célula de Horas (que o jateamento referencia)
        # arredonda, replicando a mesma perda de precisão intermediária.
        peso_caldeiraria_ref = _ref(linhas_p, "caldeiraria_peso_kg")
        peso_caldeiraria = f'IF({peso_caldeiraria_ref}="",{peso},{peso_caldeiraria_ref})'
        horas_cheio = f"({peso_caldeiraria})*{_ref(linhas_p, 'caldeiraria_fator_h_kg')}"
        return f"={horas_cheio}*{_ref(linhas_p, 'caldeiraria_valor_hora')}", f"=ROUND({horas_cheio},2)"

    if codigo == "jateamento_pintura_mo":
        if linha_caldeiraria_horas_cel is None:
            return None, None
        horas = f"{linha_caldeiraria_horas_cel}/{_ref(linhas_p, 'jateamento_divisor')}"
        return f"={horas}*{_ref(linhas_p, 'jateamento_valor_hora')}", f"=ROUND({horas},2)"

    if codigo == "solda":
        peso_consumivel_ref = _ref(linhas_p, "solda_peso_kg")
        peso_consumivel = f'IF({peso_consumivel_ref}="",{peso},{peso_consumivel_ref})'
        kg_consumivel = f"({peso_consumivel})*{_ref(linhas_p, 'solda_fator_consumo')}"
        valor_consumivel = f"({kg_consumivel})*{_ref(linhas_p, 'solda_preco_consumivel')}"

        peso_gas_ref = _ref(linhas_p, "solda_gas_peso_kg")
        peso_gas = f'IF({peso_gas_ref}="",({kg_consumivel}),{peso_gas_ref})'
        valor_gas = f"({peso_gas})*{_ref(linhas_p, 'solda_fator_gas')}*{_ref(linhas_p, 'solda_preco_gas')}"

        carbono = f"({valor_consumivel})+({valor_gas})"
        inox = (
            f"{_ref(linhas_p, 'solda_inox_qtd_kg')}*"
            f"({_ref(linhas_p, 'solda_inox_preco_consumivel')}+{_ref(linhas_p, 'solda_inox_fator_gas')}*{_ref(linhas_p, 'solda_inox_preco_gas')})"
        )
        return (
            f"=({carbono})+({inox})",
            None,
        )

    if codigo == "pintura_material":
        return (
            f"={_ref(linhas_p, 'area_pintura_m2')}*{_ref(linhas_p, 'pintura_fator_l_m2')}*"
            f"({_ref(linhas_p, 'pintura_preco_fundo')}+{_ref(linhas_p, 'pintura_preco_acabamento')})",
            None,
        )

    if codigo == "ndt":
        return f"={peso}*{_ref(linhas_p, 'ndt_valor_kg')}", None

    if codigo == "engenharia":
        return f"={_ref(linhas_p, 'qtd_posicoes_engenharia')}*{_ref(linhas_p, 'engenharia_valor_unitario')}", None

    if codigo == "embalagem":
        peso_embalagem_ref = _ref(linhas_p, "embalagem_peso_kg")
        peso_embalagem = f'IF({peso_embalagem_ref}="",{peso},{peso_embalagem_ref})'
        return f"=({peso_embalagem})*{_ref(linhas_p, 'embalagem_valor_kg')}", None

    if codigo == "transporte":
        peso_transporte_ref = _ref(linhas_p, "transporte_peso_kg")
        peso_transporte = f'IF({peso_transporte_ref}="",{peso},{peso_transporte_ref})'
        return f"=({peso_transporte})*{_ref(linhas_p, 'transporte_valor_kg')}", None

    if codigo == "energia":
        peso_energia_ref = _ref(linhas_p, "energia_peso_kg")
        peso_energia = f'IF({peso_energia_ref}="",{peso},{peso_energia_ref})'
        return f"=({peso_energia})*{_ref(linhas_p, 'energia_valor_kg')}", None

    return None, None  # usinagem, itens_padrao — sem taxa única, fica estático


def gerar_excel_orcamento(entrada: dict, resultado, params: dict | None = None) -> bytes:
    """entrada: o dict já adaptado (o mesmo que POST /orcamento espera).
    resultado: o ResultadoOrcamento já calculado (evita recalcular e
    garante que os valores estáticos batem com o que apareceu na tela)."""
    params = params or PARAMETROS_PADRAO

    wb = Workbook()
    ws = wb.active
    ws.title = "Orçamento"
    ws.sheet_view.showGridLines = False

    linhas_p = _montar_parametros(ws, entrada, params, resultado.comercial)

    # ---- Cabeçalho do relatório ----
    # Empresa que emite (Macfab) em destaque — logo + nome + CNPJ; "FC
    # Nexus" (nome do software) vira uma marca pequena e discreta embaixo,
    # não o título principal. Pedido explícito do usuário.
    r = 1
    ws.row_dimensions[r].height = 34
    if _CAMINHO_LOGO_MACFAB.exists():
        logo = ImagemExcel(str(_CAMINHO_LOGO_MACFAB))
        logo.width = 110
        logo.height = 42
        ws.add_image(logo, f"A{r}")
        _mesclar_e_estilizar(ws, r, "A", "B", "", _FONTE_TITULO, _FILL_TITULO)
        col_titulo_ini = "C"
    else:
        col_titulo_ini = "A"
    _mesclar_e_estilizar(ws, r, col_titulo_ini, "G", _MACFAB_NOME, _FONTE_TITULO, _FILL_TITULO, "left")
    r += 1
    gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M")
    if _CAMINHO_LOGO_MACFAB.exists():
        _mesclar_e_estilizar(ws, r, "A", "B", "", _FONTE_SUBTITULO, _FILL_TITULO)
    _mesclar_e_estilizar(
        ws, r, col_titulo_ini, "G", f"CNPJ {_MACFAB_CNPJ}  ·  Gerado em {gerado_em}",
        _FONTE_SUBTITULO, _FILL_TITULO, "left",
    )
    r += 1
    if _CAMINHO_LOGO_MACFAB.exists():
        _mesclar_e_estilizar(ws, r, "A", "B", "", _FONTE_MARCA_DISCRETA, _FILL_TITULO)
    _mesclar_e_estilizar(
        ws, r, col_titulo_ini, "G", "FC Nexus — Orçamento Industrial I.A.",
        _FONTE_MARCA_DISCRETA, _FILL_TITULO, "left",
    )
    r += 2

    # ---- Identificação do cliente ----
    # Pedido explícito do usuário: painel de identificação (CNPJ/nome/
    # endereço/revisão/condição de pagamento/pedido) preenchido no topo da
    # página — reflete aqui também, pra o Excel bater com o relatório PDF.
    ident = entrada.get("identificacao_cliente") or {}
    if any(ident.get(c) for c in ("cnpj", "nome", "endereco", "revisao", "condicao_pagamento", "pedido")):
        linha_ident = r

        def add_ident(rotulo: str, valor):
            nonlocal r
            ws.cell(row=r, column=1, value=rotulo).font = Font(size=9, bold=True)
            ws.cell(row=r, column=2, value=valor or "").font = Font(size=9)
            ws.merge_cells(f"B{r}:D{r}")
            r += 1

        add_ident("CNPJ", ident.get("cnpj"))
        add_ident("Cliente", ident.get("nome"))
        add_ident("Endereço", ident.get("endereco"))
        add_ident("Revisão", ident.get("revisao"))
        add_ident("Condição de pagamento", ident.get("condicao_pagamento"))
        add_ident("Pedido", ident.get("pedido"))
        for linha_b in range(linha_ident, r):
            _bordar_linha(ws, linha_b, 4)
        r += 1

    # ---- Matéria-prima ----
    itens_mp = entrada.get("materia_prima") or []
    ref_total_materia_prima = None
    if itens_mp:
        _mesclar_e_estilizar(ws, r, "A", "G", "MATÉRIA-PRIMA", _FONTE_SECAO, _FILL_SECAO)
        r += 1
        _cabecalho_tabela(ws, r, ["Descrição", "Peso (kg)", "Preço/kg", "Bruto", "ICMS", "PIS/COFINS", "Líquido"])
        r += 1
        icms, pis_cofins = ALIQUOTAS_COMPRA_POR_TIPO["materia_prima"]
        primeira_linha_mp = r
        for item in itens_mp:
            ws.cell(row=r, column=1, value=item["descricao"])
            ws.cell(row=r, column=2, value=item["peso_kg"]).number_format = "0.00"
            ws.cell(row=r, column=3, value=item["preco_kg"]).number_format = MOEDA
            ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format = MOEDA
            ws.cell(row=r, column=5, value=icms).number_format = PERCENTUAL
            ws.cell(row=r, column=6, value=pis_cofins).number_format = PERCENTUAL
            ws.cell(row=r, column=7, value=f"=D{r}*(1-E{r}-F{r})").number_format = MOEDA
            _bordar_linha(ws, r, 7)
            r += 1
        ws.cell(row=r, column=1, value="Total matéria-prima").font = _FONTE_TOTAL
        celula_total = ws.cell(row=r, column=7, value=f"=SUM(G{primeira_linha_mp}:G{r - 1})")
        celula_total.number_format = MOEDA
        celula_total.font = _FONTE_TOTAL
        _bordar_linha(ws, r, 7)
        ref_total_materia_prima = f"$G${r}"
        r += 2

    # ---- Processos ----
    _mesclar_e_estilizar(ws, r, "A", "G", "PROCESSOS / CUSTO", _FONTE_SECAO, _FILL_SECAO)
    r += 1
    _cabecalho_tabela(ws, r, ["Processo", "Horas", "Bruto", "ICMS", "PIS/COFINS", "Líquido", ""])
    r += 1
    primeira_linha_proc = r
    horas_caldeiraria_cel: str | None = None
    for linha in resultado.linhas:
        if linha.codigo == "materia_prima":
            continue

        ws.cell(row=r, column=1, value=linha.descricao)

        formula_bruto, formula_horas = _formula_bruto_processo(
            linha.codigo, entrada, linhas_p,
            horas_caldeiraria_cel if linha.codigo == "jateamento_pintura_mo" else None,
        )

        if formula_horas is not None:
            ws.cell(row=r, column=2, value=formula_horas).number_format = "0.00"
            if linha.codigo == "caldeiraria":
                horas_caldeiraria_cel = f"$B${r}"
        elif linha.horas is not None:
            ws.cell(row=r, column=2, value=linha.horas).number_format = "0.00"
            if linha.codigo == "caldeiraria":
                horas_caldeiraria_cel = f"$B${r}"

        if formula_bruto is not None:
            ws.cell(row=r, column=3, value=formula_bruto).number_format = MOEDA
        else:
            ws.cell(row=r, column=3, value=round(linha.valor_bruto, 2)).number_format = MOEDA

        ws.cell(row=r, column=4, value=linha.aliquota_icms).number_format = PERCENTUAL
        ws.cell(row=r, column=5, value=linha.aliquota_pis_cofins).number_format = PERCENTUAL
        ws.cell(row=r, column=6, value=f"=C{r}*(1-D{r}-E{r})").number_format = MOEDA
        _bordar_linha(ws, r, 7)
        r += 1

    # Pedido explícito do usuário: a matéria-prima (que já tem sua própria
    # tabela detalhada acima) também entra como uma linha aqui, pro "Total
    # processo/custo" desta seção já sair fechado (processos + matéria-
    # prima) — por isso o Resumo Comercial abaixo NÃO soma matéria-prima
    # de novo (contaria em dobro).
    if ref_total_materia_prima:
        ws.cell(row=r, column=1, value="Matéria-prima").font = Font(size=10)
        ws.cell(row=r, column=6, value=f"={ref_total_materia_prima}").number_format = MOEDA
        _bordar_linha(ws, r, 7)
        r += 1

    ws.cell(row=r, column=1, value="Total processo/custo").font = _FONTE_TOTAL
    celula_total_proc = ws.cell(row=r, column=6, value=f"=SUM(F{primeira_linha_proc}:F{r - 1})")
    celula_total_proc.number_format = MOEDA
    celula_total_proc.font = _FONTE_TOTAL
    _bordar_linha(ws, r, 7)
    ref_total_processos = f"$F${r}"
    r += 2

    # ---- Resumo comercial ----
    linha_resumo_ini = r
    _mesclar_e_estilizar(ws, r, "A", "G", "RESUMO COMERCIAL", _FONTE_SECAO, _FILL_SECAO)
    r += 1

    # ref_total_processos já inclui a matéria-prima (linha "Matéria-prima"
    # somada dentro de "Total processo/custo" acima) — não soma de novo
    # aqui, senão conta a matéria-prima em dobro no custo industrial.
    formula_custo_industrial = f"={ref_total_processos}"

    def linha_resumo(rotulo: str, formula, fmt: str):
        nonlocal r
        ws.cell(row=r, column=1, value=rotulo).font = _FONTE_RESUMO_ROTULO
        celula = ws.cell(row=r, column=3, value=formula)
        celula.font = _FONTE_RESUMO_VALOR
        celula.number_format = fmt
        return celula

    linha_peso = r
    linha_resumo("Peso líquido (kg)", f"={_ref(linhas_p, 'peso_liquido_kg')}", "0.00")
    r += 1
    linha_custo = r
    linha_resumo("Custo industrial", formula_custo_industrial, MOEDA)
    r += 1
    linha_margem = r
    linha_resumo("Fator de margem", f"={_ref(linhas_p, 'fator_margem')}", "0.00")
    r += 1
    linha_aliquota = r
    linha_resumo("Alíquota de venda", f"={_ref(linhas_p, 'aliquota_venda')}", PERCENTUAL)
    r += 1

    custo_cel, margem_cel, aliquota_cel, peso_cel = f"$C${linha_custo}", f"$C${linha_margem}", f"$C${linha_aliquota}", f"$C${linha_peso}"

    linha_venda_com = r
    linha_resumo("Preço de venda (c/ impostos)", f"={custo_cel}*(1+{margem_cel})", MOEDA)
    venda_com_cel = f"$C${linha_venda_com}"
    r += 1
    linha_imposto = r
    linha_resumo("Imposto a pagar", f"={venda_com_cel}*{aliquota_cel}", MOEDA)
    imposto_cel = f"$C${linha_imposto}"
    r += 1
    linha_resumo("Margem de lucro", f"={venda_com_cel}-{imposto_cel}-{custo_cel}", MOEDA)
    r += 1
    linha_resumo("Preço de venda (s/ impostos)", f"={venda_com_cel}*(1-{aliquota_cel})", MOEDA)
    r += 1
    linha_resumo("R$/kg", f"=IF({peso_cel}=0,0,{venda_com_cel}/{peso_cel})", MOEDA)
    r += 1

    for linha_i in range(linha_resumo_ini + 1, r):
        _bordar_linha(ws, linha_i, 3)

    r += 1
    # Preço final em destaque
    _mesclar_e_estilizar(ws, r, "A", "D", "PREÇO DE VENDA (C/ IMPOSTOS)", _FONTE_PRECO_FINAL, _FILL_PRECO_FINAL)
    ws.row_dimensions[r].height = 26
    ws.merge_cells(f"E{r}:G{r}")
    celula_preco_final = ws[f"E{r}"]
    celula_preco_final.value = f"={venda_com_cel}"
    celula_preco_final.font = _FONTE_PRECO_FINAL
    celula_preco_final.fill = _FILL_PRECO_FINAL
    celula_preco_final.number_format = MOEDA
    celula_preco_final.alignment = Alignment(horizontal="right", vertical="center")
    ultima_linha = r

    # ---- Layout A4 ----
    ws.column_dimensions["A"].width = 34
    for col in ("B", "C", "D", "E", "F", "G"):
        ws.column_dimensions[col].width = 13

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6, header=0.3, footer=0.3)
    ws.print_area = f"A1:G{ultima_linha}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

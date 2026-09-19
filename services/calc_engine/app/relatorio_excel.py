"""Excel parametrizado da BOM extraída pelo relatório técnico por IA
(services/extractor/app/ai_fallback/relatorio_tecnico.py) — vai e volta:

  exportar (gerar_excel)  -> planilha editável com uma coluna por medida
                             que o motor de geometria usa, pra revisar
                             fora do navegador.
  importar (calcular_itens_da_planilha) -> lê de volta e RECALCULA o peso
                             de cada item pelo motor determinístico
                             (app/geometria_dispatch.py) — nunca aceita
                             peso vindo da planilha/IA direto. Item sem
                             tipo de geometria reconhecido ou com medida
                             faltando vira "ignorado" em vez de forçar um
                             cálculo errado.
"""

from __future__ import annotations

import io

from app.geometria_dispatch import TIPOS_GEOMETRIA, calcular_peso
from app.materiais_catalogo import DENSIDADE_ACO_CARBONO_PADRAO, buscar_densidade

# (chave do item, rótulo da coluna) — mesma ordem na planilha de ida e de volta.
COLUNAS: list[tuple[str, str]] = [
    ("posicao", "Posição"),
    ("descricao", "Descrição"),
    ("quantidade", "Quantidade"),
    ("norma", "Norma/Material"),
    ("tipo_geometria", "Tipo de geometria"),
    ("comprimento_mm", "Comprimento (mm)"),
    ("largura_mm", "Largura (mm)"),
    ("espessura_mm", "Espessura (mm)"),
    ("diametro_mm", "Diâmetro (mm)"),
    ("diametro_externo_mm", "Diâmetro externo (mm)"),
    ("diametro_interno_mm", "Diâmetro interno (mm)"),
    ("diametro_maior_mm", "Diâmetro maior (mm)"),
    ("diametro_menor_mm", "Diâmetro menor (mm)"),
    ("base_mm", "Base (mm)"),
    ("altura_mm", "Altura (mm)"),
    ("base_menor_mm", "Base menor (mm)"),
    ("base_maior_mm", "Base maior (mm)"),
    ("diagonal_maior_mm", "Diagonal maior (mm)"),
    ("diagonal_menor_mm", "Diagonal menor (mm)"),
    ("aba_mm", "Aba (mm)"),
    ("angulo_graus", "Ângulo (graus)"),
    ("espessura_parede_mm", "Espessura da parede (mm)"),
    ("peso_kg_m", "Peso por metro — perfil (kg/m)"),
    ("confianca", "Confiança da IA (0 a 1)"),
]

_CHAVES = [chave for chave, _ in COLUNAS]


def gerar_excel(itens: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista de materiais (IA)"

    ws.append([rotulo for _, rotulo in COLUNAS])
    for celula in ws[1]:
        celula.font = celula.font.copy(bold=True)

    for item in itens:
        ws.append([item.get(chave) for chave in _CHAVES])

    ultima_linha = max(len(itens) + 1, 200)  # margem pra adicionar itens novos na mão
    col_tipo = _CHAVES.index("tipo_geometria") + 1
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(TIPOS_GEOMETRIA.keys()) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(col_tipo)}2:{get_column_letter(col_tipo)}{ultima_linha}")

    for idx, (_, rotulo) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(rotulo) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def ler_excel(conteudo: bytes) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    ws = wb.active

    itens = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() != "" for v in linha):
            continue
        valores = list(linha) + [None] * (len(_CHAVES) - len(linha))
        itens.append(dict(zip(_CHAVES, valores)))
    return itens


def _para_float(valor) -> float | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def calcular_itens_da_planilha(itens_planilha: list[dict]) -> tuple[list[dict], list[dict]]:
    """Devolve (itens_calculados, itens_ignorados). `itens_calculados` já
    no formato de ItemCalculado (ver apps/web/src/lib/types.ts) — peso
    sempre vem de `calcular_peso` (motor determinístico), nunca de um
    valor digitado na planilha."""
    calculados: list[dict] = []
    ignorados: list[dict] = []

    for i, bruto in enumerate(itens_planilha, start=1):
        posicao = str(bruto.get("posicao") or i).strip() or str(i)
        descricao = str(bruto.get("descricao") or "").strip()
        norma = (bruto.get("norma") or "").strip() or None
        quantidade = _para_float(bruto.get("quantidade")) or 1.0
        tipo = (bruto.get("tipo_geometria") or "").strip() or None

        if not tipo or tipo not in TIPOS_GEOMETRIA:
            ignorados.append({
                "posicao": posicao, "descricao": descricao,
                "motivo": "Tipo de geometria vazio ou não reconhecido — escolha um item da lista suspensa na coluna 'Tipo de geometria'.",
            })
            continue

        rotulo_tipo, campos = TIPOS_GEOMETRIA[tipo]
        medidas: dict[str, float] = {}
        faltando: list[str] = []
        for chave, rotulo_campo, _unidade in campos:
            if chave == "densidade_kg_m3":
                medidas[chave] = buscar_densidade(norma) or DENSIDADE_ACO_CARBONO_PADRAO
                continue
            valor = _para_float(bruto.get(chave))
            if valor is None:
                faltando.append(rotulo_campo)
            else:
                medidas[chave] = valor

        if faltando:
            ignorados.append({
                "posicao": posicao, "descricao": descricao,
                "motivo": f"Faltam medidas pra '{rotulo_tipo}': {', '.join(faltando)}.",
            })
            continue

        try:
            resultado = calcular_peso(tipo, medidas, quantidade)
        except Exception as exc:
            ignorados.append({"posicao": posicao, "descricao": descricao, "motivo": str(exc)})
            continue

        calculados.append({
            "posicao": posicao,
            "tipo": tipo,
            "tipoRotulo": rotulo_tipo,
            "descricao": descricao,
            "norma": norma or "",
            "quantidade": quantidade,
            "peso_kg": round(resultado.peso_kg, 3),
            "memoria_calculo": resultado.memoria,
            "formSnapshot": {k: str(v) for k, v in medidas.items()},
        })

    return calculados, ignorados

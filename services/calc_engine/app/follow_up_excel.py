"""Exportar o Follow up para Excel — pedido do usuário: botão na aba que
baixa a lista como está na tela (mesmos filtros e ordem, que o app manda
pelos ids), com todas as colunas e a 1ª foto de cada pedido dentro da
célula, como na aba "Gerencia" original. Só apresentação — nada é gravado.

Sem fotos (botão "Excel sem foto"), pedido do usuário: tira a Foto e as
etapas (ENG..PIN) e fica neutro — sem cores, só cabeçalho em negrito com
filtro e bordas finas.
"""

from __future__ import annotations

import io
from datetime import date

from app.follow_up import CAMPOS, ETAPAS, carregar_itens
from app.orcamentos_salvos import _conectar

MOEDA = '"R$" #,##0.00'
PESO = "#,##0.00"
DATA = "dd/mm/yyyy"
_MOEDAS = {"orcamento_terceirizado_unid", "orcamento_custo_macfab_unid", "preco_previsto"}
_PESOS = {"peso_unid", "peso_total"}
_DATAS = {"prazo_contratual", "coleta"}

# Mesmos rótulos da tabela do app; ordem da planilha original.
_ROTULOS = {
    "imagem": "Foto", "po": "PO", "prazo_contratual": "Prazo contratual", "cliente": "Cliente",
    "quantidade": "Qtd", "mac": "MAC", "desenho": "Desenho", "descricao": "Descrição",
    "coleta": "Coleta", "status": "Status", "fornecedor": "Fornecedor",
    "orcamento_terceirizado_unid": "Orç. terceirizado unid", "orcamento_custo_macfab_unid": "Custo Macfab unid",
    "obs_felipe_marcelo": "Obs. Felipe / Marcelo", "obs_alisson": "Obs. Alisson", "cor2": "COR2", "cor_2": "COR-2",
    "plano_pintura": "Plano de pintura", "st": "ST", "nf": "NF", "tipagem": "Tipagem", "peso_unid": "Peso unid",
    "peso_total": "Peso total", "ano": "Ano", "preco_previsto": "Preço previsto",
}
COLUNAS = [c for c, _, _ in CAMPOS if c != "d"]
_LARGURAS = {"imagem": 12, "descricao": 45, "obs_felipe_marcelo": 30, "obs_alisson": 30, "plano_pintura": 35,
             "cliente": 14, "po": 18, "mac": 14, "desenho": 18, "fornecedor": 18}

FOTO_ALTURA_PX = 60


def cor_grupo_pintura(indice: int) -> str:
    """Mesma cor do app (corGrupoPintura em apps/web/src/lib/followUp.ts), em hex
    sem '#', clareada como o tom da tabela."""
    import colorsys

    matiz = (indice * 137.508) % 360 / 360
    luz = (58, 45, 70)[indice % 3] / 100
    r, g, b = colorsys.hls_to_rgb(matiz, luz, 0.7)
    # tom suave (35% da cor sobre branco) — legível impresso
    r, g, b = (1 - 0.35 * (1 - x) for x in (r, g, b))
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _valor(item: dict, campo: str):
    if campo == "coleta":
        campo_data = item.get("coleta_data")
        return date.fromisoformat(campo_data) if campo_data else item.get("coleta")
    v = item.get(campo)
    if campo == "prazo_contratual" and v:
        return date.fromisoformat(v[:10])
    return v


def _miniatura(conteudo: bytes) -> tuple[io.BytesIO, int, int] | None:
    """PNG pequeno da foto (a original chega a vários MB) — None se não abrir."""
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(conteudo))
        img.load()
    except Exception:
        return None
    if img.mode not in ("RGB", "RGBA"):
        img = img.convert("RGBA")
    escala = FOTO_ALTURA_PX / max(1, img.height)
    largura = max(1, min(int(img.width * escala), 3 * FOTO_ALTURA_PX))
    img = img.resize((largura, FOTO_ALTURA_PX))
    saida = io.BytesIO()
    img.save(saida, format="PNG", optimize=True)
    saida.seek(0)
    return saida, largura, FOTO_ALTURA_PX


def gerar_excel(ids: list[str], com_fotos: bool = True) -> bytes:
    import openpyxl
    from openpyxl.drawing.image import Image as ImagemExcel
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    colunas = COLUNAS if com_fotos else [c for c in COLUNAS if c != "imagem" and c not in ETAPAS]

    with _conectar() as conn, conn.cursor() as cur:
        por_id = {i["id"]: i for i in carregar_itens(cur)}
        itens = [por_id[i] for i in ids if i in por_id]
        fotos: dict[str, bytes] = {}
        if com_fotos:
            shas = list({i["imagens"][0]["sha256"] for i in itens if i["imagens"]})
            if shas:
                cur.execute("select sha256, conteudo from follow_up_midias where sha256 = any(%s)", (shas,))
                fotos = {sha: bytes(c) for sha, c in cur.fetchall()}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Follow up"
    ws.append([_ROTULOS.get(c, c.upper()) for c in colunas])
    cabecalho = PatternFill("solid", fgColor="E7E5E4")
    for celula in ws[1]:
        celula.font = Font(bold=True)
        if com_fotos:
            celula.fill = cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col = {c: i for i, c in enumerate(colunas, start=1)}
    contagem_grupo: dict[int, int] = {}
    for i in itens:
        if i.get("grupo_pintura"):
            contagem_grupo[i["grupo_pintura"]] = contagem_grupo.get(i["grupo_pintura"], 0) + 1
    larguras = dict(_LARGURAS)
    miniaturas: dict[str, tuple[io.BytesIO, int, int] | None] = {}
    for linha, item in enumerate(itens, start=2):
        ws.append([None if c == "imagem" else _valor(item, c) for c in colunas])
        for c in _MOEDAS:
            ws.cell(row=linha, column=col[c]).number_format = MOEDA
        for c in _PESOS:
            ws.cell(row=linha, column=col[c]).number_format = PESO
        for c in _DATAS:
            ws.cell(row=linha, column=col[c]).number_format = DATA
        for c in ("po", "mac", "desenho", "nf", "tipagem"):
            ws.cell(row=linha, column=col[c]).number_format = "@"
        for celula in ws[linha]:
            celula.alignment = Alignment(vertical="center", wrap_text=celula.column == col["descricao"])
        if com_fotos and item.get("grupo_pintura") and contagem_grupo.get(item["grupo_pintura"], 0) >= 2:
            fundo = PatternFill("solid", fgColor=cor_grupo_pintura(item["grupo_pintura"]))
            for c in ("cor2", "cor_2", "plano_pintura"):
                ws.cell(row=linha, column=col[c]).fill = fundo

        if com_fotos and item["imagens"]:
            sha = item["imagens"][0]["sha256"]
            if sha not in miniaturas:
                miniaturas[sha] = _miniatura(fotos[sha]) if sha in fotos else None
            mini = miniaturas[sha]
            if mini:
                buf, w, h = mini
                buf.seek(0)
                img = ImagemExcel(io.BytesIO(buf.getvalue()))
                img.width, img.height = w, h
                ws.add_image(img, f"A{linha}")
                ws.row_dimensions[linha].height = h * 0.75 + 4  # px → pontos
                if w > 80:
                    larguras["imagem"] = max(larguras["imagem"], min(w / 7 + 2, 28))

    for campo, idx in col.items():
        ws.column_dimensions[get_column_letter(idx)].width = larguras.get(campo, max(10, len(_ROTULOS.get(campo, campo)) + 3))

    ultima = max(2, len(itens) + 1)
    # Etapas em barra verde, como a formatação condicional da planilha.
    for etapa in ETAPAS if com_fotos else ():
        letra = get_column_letter(col[etapa])
        ws.column_dimensions[letra].width = 7
        ws.conditional_formatting.add(
            f"{letra}2:{letra}{ultima}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="63C384"),
        )
    intervalo = f"A1:{get_column_letter(len(colunas))}{ultima}"
    ws.auto_filter.ref = intervalo
    if com_fotos:
        ws.freeze_panes = "C2"
    else:
        ws.freeze_panes = "B2"  # PO fixo
        ws.row_dimensions[1].height = 30
        fina = Side(style="thin", color="BFBFBF")
        borda = Border(left=fina, right=fina, top=fina, bottom=fina)
        for linha_celulas in ws.iter_rows(min_row=1, max_row=ultima, max_col=len(colunas)):
            for celula in linha_celulas:
                celula.border = borda

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

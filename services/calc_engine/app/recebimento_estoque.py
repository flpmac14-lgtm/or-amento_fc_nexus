"""Recebimento — espelho da aba "GERAL" (colunas A até X) de
"J:\\5 - Almoxarifado - Recebimento\\ISO_SGQ\\Recebimento\\CONTROLE DE RECEBIMENTO E ESTOQUE.xlsm"
(ver supabase/migrations/0041_recebimento_estoque.sql).

Pedido explícito do usuário (13/10/2026): sincronizar essa aba só pra
consulta da Qualidade, a cada 15 min (roda junto da tarefa da Controle de
obras, scripts/sincronizar_controle_obras.py — a máquina da empresa enxerga o
J:). A planilha é a fonte da verdade: quando o arquivo muda (hash), as linhas
são substituídas por inteiro. Mesmo formato do Material de compra
(app/material_compra.py) pra usar a mesma tela de espelho.

A aba tem 50 mil linhas formatadas, ~19,7 mil com dado; cabeçalhos repetidos
("Coluna1"…, "NRI." e "NRI.6") — o campo sai do cabeçalho + posição.
"""

from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime, timezone

from app.controle_obras import _letra, _normalizar
from app.material_compra import _limpo
from app.orcamentos_salvos import _conectar

ABA = "GERAL"
ULTIMA_COLUNA = 24  # X
# Tipo por cabeçalho (o resto é texto/código — não formatar CQ/OC como número).
_TIPOS = {"qtd.": "numero", "qtd.7": "numero", "dt receb.": "data", "seq": "codigo", "ano": "codigo"}


def _campo(cabecalho, i: int, usados: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", _normalizar(cabecalho)).strip("_") or f"col_{_letra(i + 1).lower()}"
    campo, n = base, 2
    while campo in usados:
        campo, n = f"{base}_{n}", n + 1
    usados.add(campo)
    return campo


def ler_planilha(conteudo: bytes) -> tuple[list[dict], list[tuple[int, list]]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    nome = next((a for a in wb.sheetnames if a.strip().casefold() == ABA.casefold()), None)
    if not nome:
        raise ValueError(f"Aba '{ABA}' não encontrada. Abas: {', '.join(wb.sheetnames)}")
    ws = wb[nome]
    colunas: list[dict] = []
    linhas: list[tuple[int, list]] = []
    for n, row in enumerate(ws.iter_rows(min_row=1, max_col=ULTIMA_COLUNA, values_only=True), start=1):
        if n == 1:
            usados: set[str] = set()
            colunas = [
                {"letra": _letra(i + 1), "cabecalho": str(c or "").strip() or _letra(i + 1),
                 "campo": _campo(c, i, usados), "tipo": _TIPOS.get(_normalizar(c), "texto")}
                for i, c in enumerate(list(row) + [None] * (ULTIMA_COLUNA - len(row)))
            ]
            if not any(_normalizar(c["cabecalho"]) == "nri" for c in colunas):
                raise ValueError(f"Cabeçalho da aba {ABA} mudou (não achei a coluna NRI): {[c['cabecalho'] for c in colunas]}")
            continue
        valores = [_limpo(row[i]) if i < len(row) else None for i in range(ULTIMA_COLUNA)]
        if any(v is not None for v in valores):
            linhas.append((n, valores))
    wb.close()
    return colunas, linhas


def sincronizar(conteudo: bytes, caminho: str, modificado_em: datetime | None, forcar: bool = False) -> dict:
    """Grava o espelho se o arquivo mudou (hash). Sempre registra a verificação."""
    from psycopg.types.json import Jsonb

    sha = hashlib.sha256(conteudo).hexdigest()
    agora = datetime.now(timezone.utc)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select arquivo_sha256 from recebimento_status where id = 1")
        atual = cur.fetchone()
        if atual and atual[0] == sha and not forcar:
            cur.execute(
                "update recebimento_status set ultima_verificacao_em = %s, ultimo_erro = null, ultimo_erro_em = null where id = 1",
                (agora,),
            )
            conn.commit()
            return {"alterado": False, "sha256": sha}

    colunas, linhas = ler_planilha(conteudo)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select pg_advisory_xact_lock(hashtext('recebimento_sync'))")
        cur.execute("delete from recebimento_linhas")
        with cur.copy("copy recebimento_linhas (linha_planilha, valores) from stdin") as copia:
            for n, valores in linhas:
                copia.write_row((n, Jsonb(valores)))
        cur.execute(
            """
            insert into recebimento_status
                (id, arquivo, aba, colunas, linhas, arquivo_sha256, arquivo_modificado_em,
                 ultima_alteracao_em, ultima_verificacao_em, ultimo_erro, ultimo_erro_em)
            values (1, %s, %s, %s, %s, %s, %s, %s, %s, null, null)
            on conflict (id) do update set
                arquivo = excluded.arquivo, aba = excluded.aba, colunas = excluded.colunas, linhas = excluded.linhas,
                arquivo_sha256 = excluded.arquivo_sha256, arquivo_modificado_em = excluded.arquivo_modificado_em,
                ultima_alteracao_em = excluded.ultima_alteracao_em,
                ultima_verificacao_em = excluded.ultima_verificacao_em, ultimo_erro = null, ultimo_erro_em = null
            """,
            (caminho, ABA, Jsonb(colunas), len(linhas), sha, modificado_em, agora, agora),
        )
        conn.commit()
    return {"alterado": True, "sha256": sha, "linhas": len(linhas)}


def registrar_erro(mensagem: str) -> None:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            "update recebimento_status set ultimo_erro = %s, ultimo_erro_em = now() where id = 1",
            (mensagem[:2000],),
        )
        conn.commit()


# (hash do arquivo, linhas) da última leitura — por processo (egress do Supabase).
_CACHE_LINHAS: tuple[str, list] | None = None


def listar() -> dict:
    global _CACHE_LINHAS
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select arquivo, aba, colunas, linhas, arquivo_modificado_em, ultima_alteracao_em,
                      ultima_verificacao_em, ultimo_erro, ultimo_erro_em, arquivo_sha256
               from recebimento_status where id = 1"""
        )
        st = cur.fetchone()
        sha = st[9] if st else None
        cache = _CACHE_LINHAS
        if sha and cache and cache[0] == sha:
            linhas = cache[1]
        else:
            cur.execute("select linha_planilha, valores from recebimento_linhas order by linha_planilha")
            linhas = [{"linha": r[0], "valores": r[1]} for r in cur.fetchall()]
            _CACHE_LINHAS = (sha, linhas) if sha else None

    def _iso(v):
        return v.isoformat() if v else None

    status = None
    if st:
        status = {
            "arquivo": st[0], "aba": st[1], "colunas": st[2], "linhas": st[3],
            "arquivo_modificado_em": _iso(st[4]), "ultima_alteracao_em": _iso(st[5]),
            "ultima_verificacao_em": _iso(st[6]), "ultimo_erro": st[7], "ultimo_erro_em": _iso(st[8]),
        }
    return {"status": status, "linhas": linhas}

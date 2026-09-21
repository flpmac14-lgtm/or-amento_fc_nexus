"""Planilha dos Pedidos WEIR extraídos (ver
services/extractor/app/extraction/weir_oc.py) — mesma ideia do
relatorio_pedido_andritz.py: colunas no formato que o usuário já usa na
planilha de controle de pedidos. A coluna "Referência" (equivalente à MAC
da ANDRITZ) vem sempre editada pelo usuário antes de baixar, porque a
WEIR não tem um campo assim pra detectar automaticamente. Só
apresentação — nenhum cálculo de preço/custo acontece aqui (o ajuste do
valor já foi feito na extração, ver FATOR_AJUSTE_VALOR_WEIR).
"""

from __future__ import annotations

import io

MOEDA = '"R$" #,##0.00'

COLUNAS: list[tuple[str, str]] = [
    ("item", "Item"),
    ("codigo", "Código"),
    ("valor_total", "Valor"),
    ("quantidade", "Quantidade"),
    ("descricao", "Descrição"),
    ("referencia", "Referência"),
    ("data_entrega", "Data de entrega"),
]

_CHAVES = [chave for chave, _ in COLUNAS]


def gerar_excel_pedido_weir(itens: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedidos WEIR"

    ws.append([rotulo for _, rotulo in COLUNAS])
    for celula in ws[1]:
        celula.font = celula.font.copy(bold=True)

    col_valor = _CHAVES.index("valor_total") + 1
    for linha, item in enumerate(itens, start=2):
        ws.append([item.get(chave) for chave in _CHAVES])
        ws.cell(row=linha, column=col_valor).number_format = MOEDA

    for idx, (_, rotulo) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(rotulo) + 4)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

"""Carteira da WEIR ("Follow up Report.xlsx") preenchida — pedido do usuário
(20/10/2026). A WEIR manda sempre a mesma planilha (Pedido | Item Ped |
… | Fornecedor Confirmação); o usuário escrevia à mão, na coluna
"Fornecedor Confirmação", a previsão de coleta, se já foi entregue e a NF,
ou "Não consta em nossa carteira Macfab". Aqui a planilha dele volta igual
(mesmas fontes, cores, larguras — só abre e grava com o openpyxl) com essa
coluna preenchida:

- pedido = Pedido + "-" + Item Ped com 3 dígitos (4501690746 + 000010 →
  4501690746-010), o mesmo PO do Follow up / Controle de obras;
- Follow up (coleta marcada, ST e NF que vêm da Controle de obras) e, pro
  que não está no Follow up, a própria Controle de obras (ST e NF);
- sem data de coleta marcada → "À confirmar" em vermelho (o usuário analisa);
- coluna nova "MAC" logo depois (pedido do usuário, 20/10/2026), com a cara
  das outras colunas da planilha.
"""

from __future__ import annotations

import io
import re
from datetime import date

from app.controle_obras import listar as listar_controle_obras
from app.follow_up import listar as listar_follow_up

NAO_CONSTA = "Não consta em nossa carteira Macfab"
A_CONFIRMAR = "À confirmar"
VERMELHO = "FFFF0000"


class PlanilhaInvalida(ValueError):
    pass


def _chave(po) -> str | None:
    """'4501690746-010', '4501690746-10' → '4501690746-010'."""
    m = re.fullmatch(r"\s*(\d+)\s*-\s*(\d+)\s*", str(po or ""))
    return f"{m.group(1)}-{int(m.group(2)):03d}" if m else None


def _normalizar(texto) -> str:
    import unicodedata

    s = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", s.lower())


def _data_br(iso: str | None) -> str | None:
    if not iso:
        return None
    try:
        return date.fromisoformat(iso[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return None


def _nf(v) -> str | None:
    s = str(v or "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def _texto(st: str | None, nf: str | None, coleta: str | None) -> tuple[str, bool]:
    """(texto da coluna Fornecedor Confirmação, vermelho?)"""
    st = (st or "").strip().upper()
    if st == "E":
        return ("Entregue" + (f" em {coleta}" if coleta else "") + (f" - NF {nf}" if nf else "")), False
    if st == "C":
        return "Cancelado", False
    if nf:
        return f"Faturado - NF {nf}" + (f" - Coleta {coleta}" if coleta else ""), False
    if coleta:
        return f"Previsão de coleta {coleta}", False
    return A_CONFIRMAR, True


def _indices_cabecalho(ws) -> tuple[int, int, int, int]:
    """(linha do cabeçalho, col Pedido, col Item Ped, col Fornecedor Confirmação) — 1-based."""
    for r in range(1, min(ws.max_row, 10) + 1):
        nomes = {_normalizar(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)}
        # "Fornecedor Confirmação" vem com acento estragado no arquivo da WEIR
        # ("Confirma��o") — por isso compara só o começo.
        conf = next((c for n, c in nomes.items() if n.startswith("fornecedorconfirma")), None)
        if "pedido" in nomes and "itemped" in nomes and conf:
            return r, nomes["pedido"], nomes["itemped"], conf
    raise PlanilhaInvalida(
        "Não achei as colunas Pedido, Item Ped e Fornecedor Confirmação — confira se é o Follow up Report da WEIR."
    )


def preencher(conteudo: bytes) -> tuple[bytes, dict]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    except Exception as e:  # noqa: BLE001 — arquivo que não é xlsx
        raise PlanilhaInvalida(f"Não consegui abrir a planilha ({e}).") from e
    ws = wb.active
    cab, c_ped, c_item, c_conf = _indices_cabecalho(ws)

    follow = {}
    for i in listar_follow_up()["itens"]:
        k = _chave(i.get("po"))
        if k:
            follow[k] = i
    controle = {}
    co = listar_controle_obras()
    colunas = [c["campo"] for c in ((co.get("status") or {}).get("colunas") or [])]
    if "po" in colunas:
        ip, ist = colunas.index("po"), colunas.index("st") if "st" in colunas else None
        inf = colunas.index("nf") if "nf" in colunas else None
        imac = colunas.index("mac") if "mac" in colunas else None
        for linha in co["linhas"]:
            v = linha["valores"]
            k = _chave(v[ip] if ip < len(v) else None)
            if k:
                controle[k] = {
                    "st": v[ist] if ist is not None and ist < len(v) else None,
                    "nf": v[inf] if inf is not None and inf < len(v) else None,
                    "mac": v[imac] if imac is not None and imac < len(v) else None,
                }

    # Coluna MAC do lado da Fornecedor Confirmação: cabeçalho igual ao das
    # outras colunas (cinza, Arial 9 negrito); as células copiam o estilo da
    # coluna Pedido da mesma linha (fonte, fundo amarelo do IMEDIATO…).
    from copy import copy

    c_mac = c_conf + 1
    ws.insert_cols(c_mac)
    ref = ws.cell(cab, c_ped)
    cab_mac = ws.cell(cab, c_mac, "MAC")
    for atributo in ("font", "fill", "border", "alignment", "protection"):
        setattr(cab_mac, atributo, copy(getattr(ref, atributo)))
    from openpyxl.utils import get_column_letter

    ws.column_dimensions[get_column_letter(c_mac)].width = 14

    resumo = {"linhas": 0, "entregues": 0, "faturados": 0, "com_coleta": 0, "a_confirmar": 0, "nao_consta": 0, "cancelados": 0}
    for r in range(cab + 1, ws.max_row + 1):
        pedido = str(ws.cell(r, c_ped).value or "").strip()
        item = str(ws.cell(r, c_item).value or "").strip()
        if not pedido or not item.isdigit():
            continue
        resumo["linhas"] += 1
        k = f"{pedido}-{int(item):03d}"
        celula = ws.cell(r, c_conf)
        vermelho = False
        mac = None
        if k in follow:
            mac = follow[k].get("mac")
        elif k in controle:
            mac = controle[k]["mac"]
        celula_mac = ws.cell(r, c_mac, str(mac).strip() if mac not in (None, "") else None)
        origem = ws.cell(r, c_ped)
        for atributo in ("font", "fill", "border", "alignment", "number_format", "protection"):
            setattr(celula_mac, atributo, copy(getattr(origem, atributo)))
        if k in follow:
            i = follow[k]
            texto, vermelho = _texto(i.get("st"), _nf(i.get("nf")), _data_br(i.get("coleta_data")) or (i.get("coleta") or "").strip() or None)
        elif k in controle:
            texto, vermelho = _texto(controle[k]["st"], _nf(controle[k]["nf"]), None)
        else:
            texto = NAO_CONSTA
            resumo["nao_consta"] += 1
        if texto.startswith("Entregue"):
            resumo["entregues"] += 1
        elif texto.startswith("Faturado"):
            resumo["faturados"] += 1
        elif texto.startswith("Previsão"):
            resumo["com_coleta"] += 1
        elif texto == "Cancelado":
            resumo["cancelados"] += 1
        elif texto == A_CONFIRMAR:
            resumo["a_confirmar"] += 1
        celula.value = texto
        # Mesma fonte da planilha (Arial 9); "À confirmar" em vermelho e negrito.
        if vermelho:
            celula.font = celula.font.copy(color=VERMELHO, bold=True)
        elif texto == NAO_CONSTA:
            celula.font = celula.font.copy(bold=True)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue(), resumo

"""Controle de obras — espelho da aba "OBRAS" (colunas H:AB) de
"J:\\6 - PCP\\Controle de obras.xlsm" (ver supabase/migrations/0014_controle_obras.sql).

Quem grava é o script agendado scripts/sincronizar_controle_obras.py, que
roda a cada 15 min NESTA máquina da empresa (o arquivo fica no servidor
local; o Render não enxerga o J:). O calc_engine só lê (`listar`).

Valores gravados como o Excel mostra (valor em cache das fórmulas, ex.:
"Kg tot"): datas viram "AAAA-MM-DD", números inteiros sem ".0", texto como
está. O tipo de cada coluna vai junto em `colunas` pro app formatar.
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from datetime import date, datetime, timezone

from app.orcamentos_salvos import _conectar

ABA = "OBRAS"
COLUNA_INICIAL = 8  # H
COLUNA_FINAL = 28  # AB
LINHA_CABECALHO = 2

# cabeçalho normalizado → (campo, tipo). Coluna fora daqui continua sendo
# importada (pelo intervalo H:AB), só que como texto.
_CAMPOS = {
    "po": ("po", "codigo"),
    "cl": ("cl", "texto"),
    "st": ("st", "texto"),
    "qt": ("qt", "numero"),
    "desenho": ("desenho", "codigo"),
    "rv": ("rv", "codigo"),
    "descricao": ("descricao", "texto"),
    "cor": ("cor", "texto"),
    "cor3": ("cor3", "texto"),
    "tipar": ("tipar", "codigo"),
    "mac": ("mac", "codigo"),
    "meta": ("meta", "data"),
    "obs": ("obs", "texto"),
    "pz-c": ("pz_c", "data"),
    "larg": ("larg", "numero"),
    "comp": ("comp", "numero"),
    "alt": ("alt", "numero"),
    "kg/pc": ("kg_pc", "numero"),
    "kg tot": ("kg_tot", "numero"),
    "nf": ("nf", "codigo"),
    "pp": ("ppu", "texto"),  # cabeçalho real é "PPµ" (µ some na normalização)
}


def _normalizar(texto) -> str:
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().casefold()


def _letra(col: int) -> str:
    s = ""
    while col:
        col, resto = divmod(col - 1, 26)
        s = chr(65 + resto) + s
    return s


def _valor(v):
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v if v.strip() else None
    return v


def ler_planilha(conteudo: bytes) -> tuple[list[dict], list[tuple[int, list]]]:
    """Devolve (colunas, [(linha, valores)]) da aba OBRAS, colunas H:AB."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    if ABA not in wb.sheetnames:
        raise ValueError(f"Aba '{ABA}' não encontrada. Abas: {', '.join(wb.sheetnames)}")
    ws = wb[ABA]
    colunas: list[dict] = []
    linhas: list[tuple[int, list]] = []
    for n, row in enumerate(
        ws.iter_rows(min_row=LINHA_CABECALHO, min_col=COLUNA_INICIAL, max_col=COLUNA_FINAL, values_only=True),
        start=LINHA_CABECALHO,
    ):
        if n == LINHA_CABECALHO:
            for i, cab in enumerate(row):
                campo, tipo = _CAMPOS.get(_normalizar(cab), (None, "texto"))
                colunas.append({"letra": _letra(COLUNA_INICIAL + i), "cabecalho": str(cab or "").strip(),
                                "campo": campo, "tipo": tipo})
            continue
        valores = [_valor(v) for v in row]
        if any(v is not None for v in valores):
            linhas.append((n, valores))
    wb.close()
    if not any(c["campo"] == "po" for c in colunas):
        raise ValueError("Cabeçalho da aba OBRAS mudou: não achei a coluna PO no intervalo H:AB (linha 2).")
    return colunas, linhas


def sincronizar(conteudo: bytes, caminho: str, modificado_em: datetime | None, forcar: bool = False) -> dict:
    """Grava o espelho se o arquivo mudou (hash). Sempre registra a verificação."""
    from psycopg.types.json import Jsonb

    sha = hashlib.sha256(conteudo).hexdigest()
    agora = datetime.now(timezone.utc)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select pg_advisory_xact_lock(hashtext('controle_obras_sync'))")
        cur.execute("select arquivo_sha256 from controle_obras_status where id = 1")
        atual = cur.fetchone()
        if atual and atual[0] == sha and not forcar:
            cur.execute(
                "update controle_obras_status set ultima_verificacao_em = %s, ultimo_erro = null, ultimo_erro_em = null where id = 1",
                (agora,),
            )
            conn.commit()
            return {"alterado": False, "sha256": sha}

        colunas, linhas = ler_planilha(conteudo)
        cur.execute("delete from controle_obras_linhas")
        with cur.copy("copy controle_obras_linhas (linha_planilha, valores) from stdin") as copia:
            for n, valores in linhas:
                copia.write_row((n, Jsonb(valores)))
        cur.execute(
            """
            insert into controle_obras_status
                (id, arquivo, aba, intervalo, colunas, linhas, arquivo_sha256, arquivo_modificado_em,
                 ultima_alteracao_em, ultima_verificacao_em, ultimo_erro, ultimo_erro_em)
            values (1, %s, %s, %s, %s, %s, %s, %s, %s, %s, null, null)
            on conflict (id) do update set
                arquivo = excluded.arquivo, aba = excluded.aba, intervalo = excluded.intervalo,
                colunas = excluded.colunas, linhas = excluded.linhas, arquivo_sha256 = excluded.arquivo_sha256,
                arquivo_modificado_em = excluded.arquivo_modificado_em,
                ultima_alteracao_em = excluded.ultima_alteracao_em,
                ultima_verificacao_em = excluded.ultima_verificacao_em, ultimo_erro = null, ultimo_erro_em = null
            """,
            (caminho, ABA, f"{_letra(COLUNA_INICIAL)}:{_letra(COLUNA_FINAL)}", Jsonb(colunas), len(linhas), sha,
             modificado_em, agora, agora),
        )
        conn.commit()
    return {"alterado": True, "sha256": sha, "linhas": len(linhas)}


def registrar_erro(mensagem: str) -> None:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            "update controle_obras_status set ultimo_erro = %s, ultimo_erro_em = now() where id = 1",
            (mensagem[:2000],),
        )
        conn.commit()


# (hash do arquivo, linhas) da última leitura — por processo do calc_engine.
_CACHE_LINHAS: tuple[str, list] | None = None


def _guardar_cache(sha: str | None, linhas: list) -> None:
    global _CACHE_LINHAS
    _CACHE_LINHAS = (sha, linhas) if sha else None


def listar() -> dict:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select arquivo, aba, intervalo, colunas, linhas, arquivo_modificado_em, ultima_alteracao_em,
                      ultima_verificacao_em, ultimo_erro, ultimo_erro_em
               from controle_obras_status where id = 1"""
        )
        st = cur.fetchone()
        cur.execute("select arquivo_sha256 from controle_obras_status where id = 1")
        sha = (cur.fetchone() or [None])[0]
        # Egress do Supabase: as linhas só mudam junto com o hash do arquivo —
        # a tela recarrega a cada 5 min, então reaproveita a última leitura.
        cache = _CACHE_LINHAS
        if sha and cache and cache[0] == sha:
            linhas = cache[1]
        else:
            cur.execute("select linha_planilha, valores from controle_obras_linhas order by linha_planilha")
            linhas = [{"linha": r[0], "valores": r[1]} for r in cur.fetchall()]
            _guardar_cache(sha, linhas)

    def _iso(v):
        return v.isoformat() if v else None

    status = None
    if st:
        status = {
            "arquivo": st[0], "aba": st[1], "intervalo": st[2], "colunas": st[3], "linhas": st[4],
            "arquivo_modificado_em": _iso(st[5]), "ultima_alteracao_em": _iso(st[6]),
            "ultima_verificacao_em": _iso(st[7]), "ultimo_erro": st[8], "ultimo_erro_em": _iso(st[9]),
        }
    return {"status": status, "linhas": linhas}

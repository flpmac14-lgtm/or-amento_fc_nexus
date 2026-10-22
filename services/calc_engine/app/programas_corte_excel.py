"""Excel "Programas de corte (FC Nexus).xlsx" — consulta de Compras.

Pedido do usuário (22/10/2026): Compras precisa ver o Nº do programa que o
Projeto aponta na Croqui de corte. Planilha SEPARADA, sem vínculo com o
MACLM.xlsx (o usuário preferiu não mexer no MACLM — 33 MB, vínculos externos,
consultas; o sistema nunca grava nele). Fica ao lado dele no J: e é
regravada no sync de 15 min (scripts/sincronizar_material_compra.py) só
quando os dados mudam.

Uma linha por pedido da Croqui de corte. A = PO+IT+POS e B = Nº programa
ficam nas duas primeiras colunas pra quem quiser fazer PROCV por conta
própria. Nº programa no jeito de Compras ("1418/1419"; número puro vira
número) e "N" quando o status é Sem Corte, como na coluna PJ do MACLM.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from app.orcamentos_salvos import _conectar

ABA = "Programas"
FUSO = ZoneInfo("America/Sao_Paulo")

# (cabeçalho, largura)
COLUNAS = [
    ("PO+IT+POS", 22), ("Nº programa", 14), ("Status", 16), ("Projetista", 12), ("MAC", 13),
    ("Descrição", 44), ("Desenho", 18), ("MP", 18), ("QTT", 8), ("Dt. Fazendo", 17), ("Dt. Feito", 17),
    ("Observação", 40),
]


def _programa(n_programa: str | None, status: str | None):
    if n_programa and n_programa.strip():
        partes = [p for p in re.split(r"[,;/\s]+", n_programa.strip()) if p]
        texto = "/".join(partes)
        return int(texto) if texto.isdigit() else texto
    if status == "Sem Corte":
        return "N"
    return None


def _data(v):
    return v.astimezone(FUSO).replace(tzinfo=None) if v else None


def _numero(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else float(v)


def linhas() -> list[tuple]:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select pedido, n_programa, status, projetista, mac, descricao, desenho, mp, qtt,
                      dt_fazendo, dt_feito, observacao
               from croqui_corte_itens where pedido is not null order by linha_planilha"""
        )
        rows = cur.fetchall()
    vistos: set[str] = set()
    saida = []
    for pedido, n_programa, status, projetista, mac, descricao, desenho, mp, qtt, dt_fazendo, dt_feito, obs in rows:
        chave = str(pedido).strip().upper()
        if chave in vistos:
            continue
        vistos.add(chave)
        saida.append((chave, _programa(n_programa, status), status, projetista, mac, descricao, desenho, mp,
                      _numero(qtt), _data(dt_fazendo), _data(dt_feito), obs))
    return saida


def gerar(dados: list[tuple], gerado_em: datetime) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append([c for c, _ in COLUNAS])
    for linha in dados:
        ws.append(list(linha))
    total = len(dados) + 1
    cabecalho = PatternFill("solid", fgColor="1E3A5F")
    for i, (_, largura) in enumerate(COLUNAS, start=1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = largura
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cabecalho
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col in ("J", "K"):
        for (c,) in ws[f"{col}2:{col}{total}"]:
            c.number_format = "dd/mm/yyyy hh:mm"
    for (c,) in ws[f"B2:B{total}"]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{total}"

    info = wb.create_sheet("Sobre")
    info["A1"] = "Programas de corte — gerado pelo FC Nexus"
    info["A1"].font = Font(bold=True, size=13)
    info["A3"] = "Atualizado em:"
    info["B3"] = gerado_em.replace(tzinfo=None)
    info["B3"].number_format = "dd/mm/yyyy hh:mm"
    info["A4"] = "Pedidos:"
    info["B4"] = len(dados)
    info["A6"] = "Vem da Croqui de corte do sistema (Nº do programa lançado pelo Projeto) e é regravado sozinho a cada 15 min."
    info["A7"] = "Não edite este arquivo: as alterações se perdem na próxima atualização. Deixe-o fechado pra ele poder atualizar."
    info["A8"] = "Nº programa \"N\" = pedido marcado como Sem Corte."
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar(caminho: str) -> dict:
    """Regrava o arquivo só se os dados mudaram (hash num .estado ao lado do calc_engine)."""
    dados = linhas()
    sha = hashlib.sha256(repr(dados).encode()).hexdigest()
    destino = Path(caminho)
    estado = Path(__file__).resolve().parent.parent / "programas_corte_excel.estado"
    if destino.exists() and estado.exists() and estado.read_text().strip() == sha:
        return {"alterado": False, "linhas": len(dados)}
    temporario = destino.with_name(destino.stem + ".tmp.xlsx")
    temporario.write_bytes(gerar(dados, datetime.now(FUSO)))
    # Troca atômica; se alguém estiver com o arquivo aberto no Excel, falha e tenta no próximo ciclo.
    os.replace(temporario, destino)
    estado.write_text(sha)
    return {"alterado": True, "linhas": len(dados)}

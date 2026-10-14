"""Data book — "filha" da Controle de obras (ver
supabase/migrations/0038_databook.sql), como o Follow up (app/follow_up_mae.py)
e a Croqui de corte (app/croqui_corte.py).

Pedido explícito do usuário (13/10/2026):
- trazer a aba "Geral Data book" de "J:\\6 - PCP\\Controle de databook.xlsm"
  com os dados; depois disso a Qualidade (ana, carol, pedro) preenche no app;
- colunas D até O são PROCV pelo PO (coluna C) na aba OBRAS da Controle de
  obras → fixas, travadas no app e atualizadas a cada 15 min;
- a coluna P (ST.BOOK) é o que a Qualidade preenche (+ a observação, Q);
- o que ainda não está vinculado (pedido ativo da mãe que não está aqui)
  entra sozinho e aparece no sininho.

Análise da planilha real (13/10/2026) que definiu as regras:
- 3.750 linhas, 3.714 PO (36 repetidos); 33 PO não existem mais na mãe
  (#N/D) — ficam com o que a planilha tinha, sem mãe.
- Todas as fixas batem com a mãe (PROCV = 1ª linha da OBRAS com o PO).
- ST.BOOK: lista PRONTO / FAZENDO / SEM BOOK / Vazio / Pronto Aguardando
  Coleta (validação de dados da planilha); "vazio" minúsculo = "Vazio".
- A (D.FINALIZAÇÃO) e B (AUTOR(A)) preenchidos em 748 linhas = quando e
  quem deu o book como pronto.
- 47 pedidos ST = A da mãe ainda não estavam na planilha → regra de entrada
  igual à do Follow up: ST = A.
- META (R) na planilha aponta pra linha errada (C5 na linha 4, C7 na 5…):
  aqui sai do mesmo PO, como as outras fixas.
"""

from __future__ import annotations

import re
import unicodedata
import uuid
from datetime import date, datetime, time, timezone
from zoneinfo import ZoneInfo

from app.follow_up_mae import _valor_mae
from app.orcamentos_salvos import _conectar

ABA = "Geral Data book"
FUSO = ZoneInfo("America/Sao_Paulo")

# campo do Data book → (campo da mãe em controle_obras_status.colunas, tipo).
# Coluna da planilha: D cliente (PROCV 2), E st (3), F quantidade (4),
# G desenho (5), H descrição (7), I revisão (6), J MAC (11), K Cor-1 (8),
# L Cor-2 (9), M plano (21), N prazo (14), O NF (20), R META (12).
CAMPOS_MAE: dict[str, tuple[str, str]] = {
    "cliente": ("cl", "texto"),
    "st": ("st", "texto"),
    "quantidade": ("qt", "numero"),
    "desenho": ("desenho", "texto"),
    "descricao": ("descricao", "texto"),
    "revisao": ("rv", "texto"),
    "mac": ("mac", "texto"),
    "cor1": ("cor", "texto"),
    "cor2": ("cor3", "texto"),
    "plano_pintura": ("ppu", "texto"),
    "prazo_contratual": ("pz_c", "data"),
    "nf": ("nf", "texto"),
    "meta": ("meta", "texto"),
}
FIXOS = list(CAMPOS_MAE)
# falta (migration 0039): o que está faltando pra finalizar o book — pedido do usuário.
EDITAVEIS = ["st_book", "observacao", "falta"]

# Colunas da aba (0-based): A D.FINALIZAÇÃO, B AUTOR(A), C PO, D..O fixas, P ST.BOOK, Q obs, R META.
_COLUNAS_PLANILHA = {
    "finalizado_em": 0, "finalizado_por": 1, "po": 2, "cliente": 3, "st": 4, "quantidade": 5, "desenho": 6,
    "descricao": 7, "revisao": 8, "mac": 9, "cor1": 10, "cor2": 11, "plano_pintura": 12,
    "prazo_contratual": 13, "nf": 14, "st_book": 15, "observacao": 16, "meta": 17,
}

# Validação de dados da planilha (S1:W1).
# + Rastreabilidade (pedido do usuário, 14/10/2026 — migration 0044): não faz parte
# do book, mas também se aponta — início e finalização em colunas próprias.
RASTREAB = "RASTREABILIDADE"
RASTREAB_FIM = "RASTREABILIDADE FINALIZADA"
ST_BOOK = ["PRONTO", "FAZENDO", "SEM BOOK", "Vazio", "Pronto Aguardando Coleta", RASTREAB, RASTREAB_FIM]
# Book pronto → grava D.FINALIZAÇÃO e AUTOR(A).
PRONTOS = {"PRONTO", "Pronto Aguardando Coleta"}


def _chave(texto) -> str:
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().casefold()


_ST_BOOK_POR_CHAVE = {_chave(s): s for s in ST_BOOK}


def _texto(v) -> str | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (date, time)):
        return v.strftime("%d/%m/%Y") if isinstance(v, date) else None
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    s = str(v).replace("\x00", "").strip()
    if not s or s.upper() in ("#N/A", "#N/D", "#REF!", "#VALUE!"):
        return None
    return s


def _numero(v) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else None
    except ValueError:
        return None


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def normalizar_st_book(v) -> str | None:
    s = _texto(v)
    return _ST_BOOK_POR_CHAVE.get(_chave(s), s) if s else None


# --- importação da planilha (uma vez) -----------------------------------------

def ler_planilha(caminho_ou_bytes) -> list[dict]:
    import io

    import openpyxl

    fonte = io.BytesIO(caminho_ou_bytes) if isinstance(caminho_ou_bytes, bytes) else caminho_ou_bytes
    wb = openpyxl.load_workbook(fonte, read_only=True, data_only=True)
    nome = next((a for a in wb.sheetnames if a.strip().casefold() == ABA.casefold()), None)
    if not nome:
        raise ValueError(f"Aba '{ABA}' não encontrada. Abas: {', '.join(wb.sheetnames)}")
    ws = wb[nome]
    cab = [_chave(c) for c in next(ws.iter_rows(min_row=2, max_row=2, max_col=18, values_only=True))]
    if cab[2] != "po" or cab[15] != "st.book":
        raise ValueError(f"Cabeçalho da aba {ABA} mudou (esperava PO em C e ST.BOOK em P): {cab}")
    itens = []
    for n, r in enumerate(ws.iter_rows(min_row=3, max_col=18, values_only=True), start=3):
        v = lambda c: r[_COLUNAS_PLANILHA[c]] if len(r) > _COLUNAS_PLANILHA[c] else None  # noqa: E731
        po = _texto(v("po"))
        if not po:
            continue
        fim = v("finalizado_em")
        item = {
            "po": po, "linha_planilha": n,
            "cliente": _texto(v("cliente")), "st": _texto(v("st")), "quantidade": _numero(v("quantidade")),
            "desenho": _texto(v("desenho")), "descricao": _texto(v("descricao")), "revisao": _texto(v("revisao")),
            "mac": _texto(v("mac")), "cor1": _texto(v("cor1")), "cor2": _texto(v("cor2")),
            "plano_pintura": _texto(v("plano_pintura")), "prazo_contratual": _data(v("prazo_contratual")),
            "nf": _texto(v("nf")), "meta": _texto(v("meta")),
            "st_book": normalizar_st_book(v("st_book")), "observacao": _texto(v("observacao")),
            "finalizado_em": fim.replace(tzinfo=FUSO) if isinstance(fim, datetime) else None,
            "finalizado_por": _texto(v("finalizado_por")),
        }
        itens.append(item)
    wb.close()
    return itens


_COLUNAS_ITEM = ["po", *FIXOS, "st_book", "observacao", "finalizado_em", "finalizado_por", "linha_planilha"]


def importar(caminho_ou_bytes, forcar: bool = False) -> dict:
    itens = ler_planilha(caminho_ou_bytes)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select pg_advisory_xact_lock(hashtext('databook'))")
        cur.execute("select count(*) from databook_itens")
        if cur.fetchone()[0] and not forcar:
            raise ValueError("O Data book já foi importado — use --forcar para apagar e importar de novo.")
        cur.execute("delete from databook_notificacoes")
        cur.execute("delete from databook_itens")
        cur.execute("delete from databook_status")
        with cur.copy(f"copy databook_itens ({', '.join(_COLUNAS_ITEM)}, origem) from stdin") as copia:
            for i in itens:
                copia.write_row([i[c] for c in _COLUNAS_ITEM] + ["importacao_planilha"])
        conn.commit()
    return {"importados": len(itens)}


# --- propagação mãe (Controle de obras) → filha --------------------------------

def _igual(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, (int, float)) or hasattr(a, "is_finite") or isinstance(b, (int, float)):
        try:
            return abs(float(a) - float(b)) < 1e-9
        except (TypeError, ValueError):
            return False
    return str(a) == str(b)


def propagar(forcar: bool = False) -> dict:
    """Leva as fixas da mãe pros itens e cria os pedidos ST = A que ainda não
    estão no Data book (no fim). Sininho: novo / encerrado (saiu de A) / reaberto."""
    agora = datetime.now(timezone.utc)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select pg_advisory_xact_lock(hashtext('databook'))")
        cur.execute("select count(*) from databook_itens")
        itens_antes = cur.fetchone()[0]
        if not itens_antes:
            return {"atualizados": 0, "novos": 0, "motivo": "Data book ainda não importado"}
        cur.execute("select colunas, arquivo_sha256 from controle_obras_status where id = 1")
        mae_st = cur.fetchone()
        if not mae_st:
            return {"atualizados": 0, "novos": 0, "motivo": "Controle de obras ainda não sincronizado"}
        # Egress do Supabase (igual 0025): só relê a mãe se ela mudou ou se
        # entrou/saiu item da filha desde a última propagação.
        cur.execute("select propagado_sha256, propagado_itens from databook_status where id = 1")
        prop = cur.fetchone()
        if not forcar and prop and prop[0] == mae_st[1] and prop[1] == itens_antes:
            return {"atualizados": 0, "novos": 0, "motivo": "mãe sem alteração"}
        indice = {c["campo"]: i for i, c in enumerate(mae_st[0]) if c.get("campo")}
        faltando = [m for m, _ in CAMPOS_MAE.values() if m not in indice] + ([] if "po" in indice else ["po"])
        if faltando:
            raise ValueError(f"Colunas da Controle de obras não encontradas: {', '.join(faltando)}")

        # PROCV exato: primeira linha com o PO (sem diferenciar maiúsculas).
        cur.execute("select linha_planilha, valores from controle_obras_linhas order by linha_planilha")
        mae: dict[str, tuple[int, list]] = {}
        for linha, valores in cur.fetchall():
            po = valores[indice["po"]]
            if po is None or str(po).strip() == "":
                continue
            mae.setdefault(_valor_mae(po, "texto").upper(), (linha, valores))

        def da_mae(valores):
            return {c: _valor_mae(valores[indice[m]], t) for c, (m, t) in CAMPOS_MAE.items()}

        cur.execute(f"select id, po, {', '.join(FIXOS)} from databook_itens")
        existentes = cur.fetchall()
        pos_existentes = {str(r[1]).upper() for r in existentes}
        atualizacoes, notificacoes = [], []
        for row in existentes:
            item_id, po, atuais = row[0], row[1], dict(zip(FIXOS, row[2:]))
            achado = mae.get(str(po).upper())
            if not achado:
                continue
            linha, valores = achado
            novos = da_mae(valores)
            if all(_igual(atuais[c], novos[c]) for c in FIXOS):
                continue
            atualizacoes.append([novos[c] for c in FIXOS] + [linha, agora, item_id])
            if (atuais["st"] == "A") != (novos["st"] == "A"):
                notificacoes.append(("encerrado" if atuais["st"] == "A" else "reaberto", item_id, po,
                                     novos["cliente"], novos["descricao"], atuais["st"], novos["st"]))
        if atualizacoes:
            cur.executemany(
                f"update databook_itens set {', '.join(f'{c} = %s' for c in FIXOS)}, mae_linha = %s, "
                "mae_sincronizada_em = %s where id = %s",
                atualizacoes,
            )

        cur.execute("select coalesce(max(linha_planilha), 2) from databook_itens")
        proxima = cur.fetchone()[0] + 1
        inseridos = []
        for po_upper, (linha, valores) in mae.items():  # dict mantém a ordem da planilha
            if po_upper in pos_existentes or _valor_mae(valores[indice["st"]], "texto") != "A":
                continue
            novos = da_mae(valores)
            novo_id = uuid.uuid4()
            po = _valor_mae(valores[indice["po"]], "texto")
            inseridos.append([novo_id, po] + [novos[c] for c in FIXOS] + [proxima, linha, agora])
            notificacoes.append(("novo", novo_id, po, novos["cliente"], novos["descricao"], None, "A"))
            proxima += 1
        if inseridos:
            cur.executemany(
                f"""insert into databook_itens (id, po, {', '.join(FIXOS)}, linha_planilha, mae_linha,
                        mae_sincronizada_em, origem)
                    values (%s, %s, {', '.join(['%s'] * len(FIXOS))}, %s, %s, %s, 'controle_obras')""",
                inseridos,
            )
        if notificacoes:
            cur.executemany(
                """insert into databook_notificacoes (tipo, item_id, po, cliente, descricao, st_antes, st_depois)
                   values (%s, %s, %s, %s, %s, %s, %s)""",
                notificacoes,
            )
        cur.execute(
            """insert into databook_status (id, propagado_sha256, propagado_itens) values (1, %s, %s)
               on conflict (id) do update set propagado_sha256 = excluded.propagado_sha256,
                 propagado_itens = excluded.propagado_itens""",
            (mae_st[1], itens_antes + len(inseridos)),
        )
        conn.commit()
    return {"atualizados": len(atualizacoes), "novos": len(inseridos),
            "encerrados": sum(1 for n in notificacoes if n[0] == "encerrado")}


# --- leitura / edição no app --------------------------------------------------

_COLUNAS_LISTA = ["id", "po", *FIXOS, *EDITAVEIS, "fazendo_em", "fazendo_por", "finalizado_em", "finalizado_por",
                  "rastreab_inicio_em", "rastreab_inicio_por", "rastreab_fim_em", "rastreab_fim_por",
                  "linha_planilha", "origem", "editado_em", "editado_por"]


def _json(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if hasattr(v, "is_finite"):  # Decimal
        f = float(v)
        return int(f) if f.is_integer() else f
    return v


def _carregar(cur, ids: list[str] | None = None) -> list[dict]:
    filtro, params = ("where id = any(%s::uuid[])", (ids,)) if ids is not None else ("", ())
    cur.execute(f"select {', '.join(_COLUNAS_LISTA)} from databook_itens {filtro} order by linha_planilha, po",
                params)
    itens = [{c: _json(v) for c, v in zip(_COLUNAS_LISTA, row)} for row in cur.fetchall()]
    for i in itens:
        i["id"] = str(i["id"])
    return itens


def _fotos_do_follow_up(cur) -> tuple[dict[str, str], dict[str, str]]:
    """1ª foto de cada pedido do Follow up, por PO e por desenho (pedido do
    usuário: a coluna Foto do Data book vem do Follow up pelo PO; sem foto
    no PO, a de outro pedido com o mesmo desenho)."""
    cur.execute(
        """select distinct on (f.id) upper(f.po), upper(f.desenho), i.midia_sha256
           from follow_up_imagens i join follow_up_itens f on f.id = i.item_id
           order by f.id, i.ordem"""
    )
    por_po: dict[str, str] = {}
    por_desenho: dict[str, str] = {}
    for po, desenho, sha in cur.fetchall():
        if po:
            por_po.setdefault(po, sha)
        if desenho:
            por_desenho.setdefault(desenho, sha)
    return por_po, por_desenho


def listar() -> dict:
    with _conectar() as conn, conn.cursor() as cur:
        itens = _carregar(cur)
        por_po, por_desenho = _fotos_do_follow_up(cur)
    for i in itens:
        sha = por_po.get(str(i["po"]).upper())
        i["foto_por"] = "po" if sha else None
        if not sha and i["desenho"]:
            sha = por_desenho.get(str(i["desenho"]).upper())
            i["foto_por"] = "desenho" if sha else None
        i["foto"] = sha
    return {"itens": itens, "st_book_opcoes": ST_BOOK}


def listar_notificacoes(limite: int = 200) -> list[dict]:
    """Sininho (mais recentes primeiro) — mesmo formato do Follow up."""
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select id, tipo, item_id, po, cliente, descricao, st_antes, st_depois, criado_em
               from databook_notificacoes order by criado_em desc, id desc limit %s""",
            (max(1, min(limite, 1000)),),
        )
        return [
            {"id": r[0], "tipo": r[1], "item_id": str(r[2]) if r[2] else None, "po": r[3], "cliente": r[4],
             "descricao": r[5], "st_antes": r[6], "st_depois": r[7], "criado_em": r[8].isoformat()}
            for r in cur.fetchall()
        ]


class CampoNaoEditavel(ValueError):
    pass


def _normalizar_alteracoes(alteracoes: dict) -> dict:
    normal = {}
    for campo, valor in alteracoes.items():
        if campo not in EDITAVEIS:
            raise CampoNaoEditavel(f"'{campo}' vem da Controle de obras (PROCV) e não é editável.")
        if campo == "st_book":
            s = normalizar_st_book(valor)
            if s is not None and s not in ST_BOOK:
                raise ValueError(f"ST.BOOK inválido: '{s}'. Use {', '.join(ST_BOOK)}.")
            normal[campo] = s
        else:
            normal[campo] = _texto(valor)
    if not normal:
        raise ValueError("Nada para salvar.")
    return normal


def _salvar(cur, ids: list[str], normal: dict, editado_por: str | None) -> None:
    agora = datetime.now(timezone.utc)
    cur.execute(
        f"select id, finalizado_em, finalizado_por, {', '.join(EDITAVEIS)} from databook_itens "
        "where id = any(%s::uuid[]) for update",
        (ids,),
    )
    for item_id, fim_em, fim_por, *valores in cur.fetchall():
        atuais = dict(zip(EDITAVEIS, valores))
        sets = {c: v for c, v in normal.items() if v != atuais[c]}
        if not sets:
            continue
        if "st_book" in sets:
            novo, st_atual = sets["st_book"], atuais["st_book"]
            # Como o projeto (Croqui de corte): FAZENDO grava o início do book.
            if novo == "FAZENDO":
                sets.update(fazendo_em=agora, fazendo_por=editado_por)
            # Rastreabilidade: marcar grava o início (e recomeça); finalizada grava o fim.
            if novo == RASTREAB:
                sets.update(rastreab_inicio_em=agora, rastreab_inicio_por=editado_por,
                            rastreab_fim_em=None, rastreab_fim_por=None)
            elif novo == RASTREAB_FIM:
                sets.update(rastreab_fim_em=agora, rastreab_fim_por=editado_por)
            # Book pronto: grava D.FINALIZAÇÃO e AUTOR(A); voltou pra não pronto: apaga.
            if novo in PRONTOS and st_atual not in PRONTOS:
                sets.update(finalizado_em=agora, finalizado_por=editado_por)
            elif novo not in PRONTOS and st_atual in PRONTOS:
                sets.update(finalizado_em=None, finalizado_por=None)
                # Guarda a finalização apagada: "Excluir status" volta com ela (pedido do usuário).
                if fim_em:
                    cur.execute(
                        "insert into databook_historico (item_id, campo, de, por, em) values (%s, 'finalizado_apagado', %s, %s, %s)",
                        (item_id, f"{fim_em.isoformat()}|{fim_por or ''}", editado_por, agora),
                    )
        # Registro de cada mudança (quem, quando, de → para) — pedido do usuário.
        cur.executemany(
            "insert into databook_historico (item_id, campo, de, para, por, em) values (%s, %s, %s, %s, %s, %s)",
            [(item_id, c, atuais[c], normal[c], editado_por, agora) for c in normal if normal[c] != atuais[c]],
        )
        cur.execute(
            f"update databook_itens set {', '.join(f'{c} = %s' for c in sets)}, editado_em = %s, editado_por = %s "
            "where id = %s",
            (*sets.values(), agora, editado_por, item_id),
        )


def excluir_status(item_id: str, por: str | None) -> dict | None:
    """Exclui o ST.BOOK marcado errado (pedido do usuário): desfaz a última troca
    de ST.BOOK ainda não desfeita — volta pro status anterior — e recalcula o
    início (último FAZENDO) e a finalização (último pronto) pelo que sobrou do
    histórico. Sem troca feita no app pra desfazer, só apaga o status."""
    agora = datetime.now(timezone.utc)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select st_book from databook_itens where id = %s for update", (item_id,))
        linha = cur.fetchone()
        if not linha:
            return None
        atual = linha[0]
        if atual is None:
            raise ValueError("Esse pedido não tem ST.BOOK pra excluir.")
        cur.execute(
            """select id, de from databook_historico
               where item_id = %s and campo = 'st_book' and not desfeito order by em desc, id desc limit 1""",
            (item_id,),
        )
        ultima = cur.fetchone()
        anterior = ultima[1] if ultima else None
        if ultima:
            cur.execute("update databook_historico set desfeito = true where id = %s", (ultima[0],))
        # O que sobrou do histórico define as datas.
        cur.execute(
            """select para, por, em from databook_historico
               where item_id = %s and campo = 'st_book' and not desfeito order by em, id""",
            (item_id,),
        )
        trocas = cur.fetchall()
        fazendo = next(((p, e) for s, p, e in reversed(trocas) if s == "FAZENDO"), None)
        pronto = next(((p, e) for s, p, e in reversed(trocas) if s in PRONTOS), None)
        r_ini = next(((p, e) for s, p, e in reversed(trocas) if s == RASTREAB), None)
        r_fim = next(((p, e) for s, p, e in reversed(trocas)
                      if s == RASTREAB_FIM and (not r_ini or e >= r_ini[1])), None)
        sets: dict = {"st_book": anterior}
        if ultima:
            # Sem FAZENDO no histórico que sobrou: o início some (veio do que foi desfeito).
            sets.update(fazendo_em=fazendo[1] if fazendo else None, fazendo_por=fazendo[0] if fazendo else None)
            sets.update(rastreab_inicio_em=r_ini[1] if r_ini else None, rastreab_inicio_por=r_ini[0] if r_ini else None,
                        rastreab_fim_em=r_fim[1] if r_fim else None, rastreab_fim_por=r_fim[0] if r_fim else None)
        if anterior not in PRONTOS:
            sets.update(finalizado_em=None, finalizado_por=None)
        elif pronto and ultima:
            sets.update(finalizado_em=pronto[1], finalizado_por=pronto[0])
        elif ultima:
            # Pronto que veio da planilha: volta a D.FINALIZAÇÃO/AUTOR que a troca errada apagou.
            cur.execute(
                """select de from databook_historico where item_id = %s and campo = 'finalizado_apagado'
                   order by em desc, id desc limit 1""",
                (item_id,),
            )
            apagado = cur.fetchone()
            if apagado and apagado[0]:
                em_txt, _, por_txt = apagado[0].partition("|")
                sets.update(finalizado_em=datetime.fromisoformat(em_txt), finalizado_por=por_txt or None)
        cur.execute(
            f"update databook_itens set {', '.join(f'{c} = %s' for c in sets)}, editado_em = %s, editado_por = %s "
            "where id = %s",
            (*sets.values(), agora, por, item_id),
        )
        cur.execute(
            "insert into databook_historico (item_id, campo, de, para, por, em) values (%s, 'st_book_excluido', %s, %s, %s, %s)",
            (item_id, atual, anterior, por, agora),
        )
        conn.commit()
        itens = _carregar(cur, [item_id])
    return itens[0] if itens else None


def historico(item_id: str) -> list[dict]:
    """Registro do pedido no Data book (mais recente primeiro)."""
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select campo, de, para, por, em, desfeito from databook_historico
               where item_id = %s and campo <> 'finalizado_apagado' order by em desc, id desc limit 200""",
            (item_id,),
        )
        return [{"campo": r[0], "de": r[1], "para": r[2], "por": r[3], "em": r[4].isoformat(), "desfeito": r[5]}
                for r in cur.fetchall()]


def editar_lote(ids: list[str], alteracoes: dict, editado_por: str | None = None) -> list[dict]:
    """Salva os editáveis numa ou em várias linhas (mesmo valor — "puxar" como no Excel)."""
    normal = _normalizar_alteracoes(alteracoes)
    ids = [str(i) for i in ids][:5000]
    if not ids:
        return []
    with _conectar() as conn, conn.cursor() as cur:
        _salvar(cur, ids, normal, editado_por)
        conn.commit()
        return _carregar(cur, ids)

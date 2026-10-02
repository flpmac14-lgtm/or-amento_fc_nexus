"""Material de compra — espelho da aba "MACLM" de
"J:\\6 - PCP\\PCP-CP\\MACLM.xlsx" (ver supabase/migrations/0020_material_compra.sql).

Pedido explícito do usuário: aba nova no módulo Follow up, planilha "mãe" de
um projeto novo (como a Controle de obras), atualizada a cada 15 min pelo
script agendado scripts/sincronizar_material_compra.py (roda na máquina da
empresa que enxerga o J:). Só as colunas pedidas, NA ORDEM pedida — achadas
pelo cabeçalho (linha 1), não pela letra, pra aguentar coluna inserida.

A aba tem duas colunas "QT" (G = quantidade do item, U = quantidade da
posição): a 2ª vira "QT_1", como o usuário chamou. Erros de fórmula
("#N/A"...) viram vazio.
"""

from __future__ import annotations

import hashlib
import io
from datetime import datetime, time, timezone

from app.controle_obras import _letra, _normalizar, _valor
from app.orcamentos_salvos import _conectar

ABA = "MACLM"
LINHA_CABECALHO = 1

# (cabeçalho na planilha, qual ocorrência dele, campo, rótulo no app, tipo)
COLUNAS: list[tuple[str, int, str, str, str]] = [
    ("PO+IT+POS", 1, "po_it_pos", "PO+IT+POS", "codigo"),
    ("PZC", 1, "pzc", "PZC", "data"),
    ("DESCRIÇÃO", 1, "descricao", "DESCRIÇÃO", "texto"),
    ("MAC1", 1, "mac1", "MAC1", "codigo"),
    ("QT", 1, "qt", "QT", "numero"),
    ("DESENHO", 1, "desenho", "DESENHO", "codigo"),
    ("ST", 1, "st", "ST", "texto"),
    ("CL", 1, "cl", "CL", "texto"),
    ("PO", 1, "po", "PO", "codigo"),
    ("Q", 1, "q", "Q", "numero"),
    ("DesenhoCJ", 1, "desenho_cj", "DesenhoCJ", "codigo"),
    ("POS", 1, "pos", "POS", "codigo"),
    ("QT", 2, "qt_1", "QT_1", "numero"),
    ("QTT", 1, "qtt", "QTT", "numero"),
    ("UN", 1, "un", "UN", "texto"),
    ("CODIGO", 1, "codigo", "CODIGO", "codigo"),
    ("MPD", 1, "mpd", "MPD", "texto"),
    ("MP", 1, "mp", "MP", "texto"),
    ("LARG", 1, "larg", "LARG", "numero"),
    ("COMP", 1, "comp", "COMP", "numero"),
    ("DI", 1, "di", "DI", "numero"),
    ("L", 1, "l", "L", "texto"),
]


def _limpo(v):
    v = _valor(v)
    if isinstance(v, time):  # célula formatada como hora (ex.: DESENHO "00:00") — 00:00 = vazio
        return None if v == time(0, 0) else v.strftime("%H:%M")
    if isinstance(v, str):
        v = v.replace("\x00", "").strip() or None  # há célula com caractere nulo — o Postgres recusa
    if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")) and len(v) <= 8:
        return None  # #N/A, #REF!, #VALUE!, #NAME?, #DIV/0!...
    return v


def _achar_colunas(cabecalho: tuple) -> list[int]:
    """Índice (0-based) de cada coluna pedida, pelo nome e ocorrência."""
    vistos: dict[str, list[int]] = {}
    for i, cab in enumerate(cabecalho):
        vistos.setdefault(_normalizar(cab), []).append(i)
    indices, faltando = [], []
    for cab, ocorrencia, *_ in COLUNAS:
        achados = vistos.get(_normalizar(cab), [])
        if len(achados) < ocorrencia:
            faltando.append(cab if ocorrencia == 1 else f"{cab} ({ocorrencia}ª)")
        else:
            indices.append(achados[ocorrencia - 1])
    if faltando:
        raise ValueError(f"Colunas não encontradas na aba {ABA} (linha 1): {', '.join(faltando)}")
    return indices


def ler_planilha(conteudo: bytes) -> tuple[list[dict], list[tuple[int, list]]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    if ABA not in wb.sheetnames:
        raise ValueError(f"Aba '{ABA}' não encontrada. Abas: {', '.join(wb.sheetnames)}")
    ws = wb[ABA]
    indices: list[int] = []
    colunas: list[dict] = []
    linhas: list[tuple[int, list]] = []
    for n, row in enumerate(ws.iter_rows(min_row=LINHA_CABECALHO, values_only=True), start=LINHA_CABECALHO):
        if n == LINHA_CABECALHO:
            indices = _achar_colunas(row)
            colunas = [{"letra": _letra(i + 1), "cabecalho": rotulo, "campo": campo, "tipo": tipo}
                       for i, (_, _, campo, rotulo, tipo) in zip(indices, COLUNAS)]
            continue
        valores = [_limpo(row[i]) if i < len(row) else None for i in indices]
        if any(v is not None for v in valores):
            linhas.append((n, valores))
    wb.close()
    return colunas, linhas


def sincronizar(conteudo: bytes, caminho: str, modificado_em: datetime | None, forcar: bool = False) -> dict:
    """Grava o espelho se o arquivo mudou (hash). Sempre registra a verificação."""
    from psycopg.types.json import Jsonb

    sha = hashlib.sha256(conteudo).hexdigest()
    agora = datetime.now(timezone.utc)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select arquivo_sha256 from material_compra_status where id = 1")
        atual = cur.fetchone()
        if atual and atual[0] == sha and not forcar:
            cur.execute(
                "update material_compra_status set ultima_verificacao_em = %s, ultimo_erro = null, ultimo_erro_em = null where id = 1",
                (agora,),
            )
            conn.commit()
            return {"alterado": False, "sha256": sha}

    # Leitura (~1 min pra 57 mil linhas) fora da transação.
    colunas, linhas = ler_planilha(conteudo)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select pg_advisory_xact_lock(hashtext('material_compra_sync'))")
        cur.execute("delete from material_compra_linhas")
        with cur.copy("copy material_compra_linhas (linha_planilha, valores) from stdin") as copia:
            for n, valores in linhas:
                copia.write_row((n, Jsonb(valores)))
        cur.execute(
            """
            insert into material_compra_status
                (id, arquivo, aba, colunas, linhas, arquivo_sha256, arquivo_modificado_em,
                 ultima_alteracao_em, ultima_verificacao_em, ultimo_erro, ultimo_erro_em)
            values (1, %s, %s, %s, %s, %s, %s, %s, %s, null, null)
            on conflict (id) do update set
                arquivo = excluded.arquivo, aba = excluded.aba, colunas = excluded.colunas, linhas = excluded.linhas,
                arquivo_sha256 = excluded.arquivo_sha256, arquivo_modificado_em = excluded.arquivo_modificado_em,
                ultima_alteracao_em = excluded.ultima_alteracao_em,
                ultima_verificacao_em = excluded.ultima_verificacao_em, ultimo_erro = null, ultimo_erro_em = null
            """,
            (caminho, ABA, Jsonb(colunas), len(linhas), sha, modificado_em, agora, agora),
        )
        conn.commit()
    return {"alterado": True, "sha256": sha, "linhas": len(linhas)}


def registrar_erro(mensagem: str) -> None:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            "update material_compra_status set ultimo_erro = %s, ultimo_erro_em = now() where id = 1",
            (mensagem[:2000],),
        )
        conn.commit()


# (hash do arquivo, linhas) da última leitura — por processo do calc_engine.
_CACHE_LINHAS: tuple[str, list] | None = None


def _guardar_cache(sha: str | None, linhas: list) -> None:
    global _CACHE_LINHAS
    _CACHE_LINHAS = (sha, linhas) if sha else None


def listar() -> dict:
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """select arquivo, aba, colunas, linhas, arquivo_modificado_em, ultima_alteracao_em,
                      ultima_verificacao_em, ultimo_erro, ultimo_erro_em
               from material_compra_status where id = 1"""
        )
        st = cur.fetchone()
        cur.execute("select arquivo_sha256 from material_compra_status where id = 1")
        sha = (cur.fetchone() or [None])[0]
        # Egress do Supabase: as linhas só mudam junto com o hash do arquivo —
        # a tela recarrega a cada 5 min, então reaproveita a última leitura.
        cache = _CACHE_LINHAS
        if sha and cache and cache[0] == sha:
            linhas = cache[1]
        else:
            cur.execute("select linha_planilha, valores from material_compra_linhas order by linha_planilha")
            linhas = [{"linha": r[0], "valores": r[1]} for r in cur.fetchall()]
            _guardar_cache(sha, linhas)

    def _iso(v):
        return v.isoformat() if v else None

    status = None
    if st:
        status = {
            "arquivo": st[0], "aba": st[1], "colunas": st[2], "linhas": st[3],
            "arquivo_modificado_em": _iso(st[4]), "ultima_alteracao_em": _iso(st[5]),
            "ultima_verificacao_em": _iso(st[6]), "ultimo_erro": st[7], "ultimo_erro_em": _iso(st[8]),
        }
    return {"status": status, "linhas": linhas}

"""Fornecedores — espelho da planilha "J:\\6 - PCP\\PCP-Fornecedor\\FORNECEDOR.xlsx"
(tabela fornecedores_planilha, supabase/migrations/0062_fornecedores.sql).

Pedido do usuário (28/10/2026): aba "Fornecedores" pro almoxarifado (leandro),
organizada por grupos parecidos e fácil de pesquisar, com telefone, e-mail,
contato e endereço. Roda no mesmo ciclo de 15 min do Material de compra
(scripts/sincronizar_material_compra.py) e só regrava quando o arquivo muda.

Abas usadas:
  FN = cadastro (B = códigos de grupo separados por "*", C = fornecedor,
       D/E = fones, F/G = celulares, H = e-mail de vendas, I = contato,
       J = marcas/fabricantes, K = materiais, N = endereço, O = cidade,
       P = CEP, Q = UF, R = CNPJ). Achadas pela posição: a linha 1 da FN
       não tem cabeçalho em todas as colunas.
  FM = nome de cada código ("ABS" = ABRASIVO). Código composto
       "CLA*CLC*CLI" = "CHAPA ALUMINIO/CARBONO/INOX" vira um nome por código.
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from datetime import datetime, timezone

from app.orcamentos_salvos import _conectar

# Códigos usados no cadastro que não têm nome na aba FM — só os óbvios pelo
# padrão dos outros (CH + C/I, FL = flange, SVTP = transporte…). O resto
# aparece como "Código XXX".
NOMES_EXTRAS = {
    "CHC": "CHAPA CARBONO", "CHI": "CHAPA INOX", "FLC": "FLANGE CARBONO", "FLI": "FLANGE INOX",
    "TQC": "TUBO QUADRADO CARBONO", "PFDC": "PERFIL DOBRADO CARBONO", "CNXT": "CONEXÕES TUBULARES",
    "SVTP": "SERVIÇO DE TRANSPORTE", "SVTPF": "SERVIÇO DE TRANSPORTE FRACIONADO", "FRETE": "FRETE",
    "FRETE-TRUCK": "FRETE (TRUCK)", "GAS": "GASES", "IF": "INFORMÁTICA", "HT": "HOTEL", "NYLON": "NYLON",
    "UHMW": "UHMW", "TECNIL": "TECNIL", "BRONZE": "BRONZE", "SDCSF": "SOLDA-CONSUMÍVEIS (FABRICANTE)",
    "ACA": "CABO DE AÇO", "AOR": "ANEL O'RING", "EMB.": "EMBALAGEM", "LBF": "LUBRIFICANTES", "LBR": "LUBRIFICANTES",
    "SVCQ": "SERVIÇO DE CONTROLE DE QUALIDADE", "SVMD": "SERVIÇO MÉDICO / CLÍNICA", "SVLE": "SERVIÇO EM EMPILHADEIRA",
    "SVLC": "SERVIÇO DE LOCAÇÃO", "SVCS": "SERVIÇO DE CONSULTORIA", "BRAL": "BARRA ALUMÍNIO", "CLAL": "CHAPA ALUMÍNIO",
    "CTAL": "CANTONEIRA ALUMÍNIO", "IMP": "IMPORTAÇÃO", "SVRF": "SERVIÇO DE REFRIGERAÇÃO / AR-CONDICIONADO",
}

# Grupos "parecidos" (o que aparece em botões na tela), pelo nome do código.
# A ordem importa: a 1ª regra que casar vence.
GRUPOS: list[tuple[str, tuple[str, ...]]] = [
    ("Serviços", ("SERVICO", "SERMOTOR")),
    ("Transporte e frete", ("TRANSPORT", "FRETE", "LOGIST")),
    ("Chapas, bobinas e telas", ("CHAPA", "BOBINA", "HARDOX", "SLITTER", "TELA")),
    ("Tubos", ("TUBO", "FEIXE TUBULAR")),
    ("Barras, perfis e vigas", ("BARRA", "VIGA", "CANTONEIRA", "PERFIL DOBRADO", "METALON", "TRILHO", "FORJADO", "ACOS ESPECIAIS", "ACO CONSTRUCAO")),
    ("Fixação (parafusos, arruelas, anéis)", ("FIXACAO", "ARRUELA", "ANEL", "CHAVETA", "PINO", "INSERTO")),
    ("Solda e gases", ("SOLDA", "GASES")),
    ("Tintas", ("TINTA",)),
    ("Jato, laser, plasma e oxicorte", ("JATO", "GRANALHA", "LASER", "PLASMA", "OXICORTE", "MICRO ESFERA", "OXIDO DE ALUMINIO")),
    ("Ferramentas e medição", ("FERRAMENTA", "FERAMENTA", "ABRASIVO", "INSTRUMENTO", "MEDIDOR", "MANOMETRO", "TERMOMETRO", "REGUA", "ANALIZADOR", "MARCADOR")),
    ("Elétrica e automação", ("ELETRIC", "ELETRON", "CABO ELETRICO", "DISJUNTOR", "CONTATOR", "MOTOR", "INVERSOR", "CLP", "SENSOR", "AUTOMA", "FUSIVEL", "RELE", "TRANSFORMADOR", "BATERIA", "SINALIZADOR", "TERMINAL", "BOTAO", "PRESSOSTATO", "TERMOSTATO", "TRASMISSOR", "INDUZIDO", "ESCOVA DE CARVAO")),
    ("Hidráulica, pneumática e conexões", ("HIDRAUL", "PNEUMAT", "VALVULA", "CONEX", "BOMBA", "MANGUEIRA", "FILTRO", "COMPRESSOR", "FLANGE", "BICO", "PURGADOR", "SECADOR", "SEPARADOR", "RESERVATORIO", "SERPENTINA", "TROCADOR")),
    ("Rolamentos, transmissão e vedação", ("ROLAMENTO", "MANCAL", "CORREIA", "CORRENTE", "ENGRENAGEM", "POLIA", "REDUTOR", "ACOPLAMENTO", "RETENTOR", "VEDA", "GAXETA", "JUNTA", "LUBRIF", "FREIO", "FUSO", "MOLA", "MOVIMENTACAO", "RODIZIO", "COXIM", "SELO")),
    ("Plásticos e borrachas", ("BORRACHA", "PLASTIC", "POLIMERO", "NYLON", "UHMW", "TECNIL", "ACRILICO", "POLICARBONATO", "ESPUMA", "FIBRA", "FELTRO", "CORTICA", "TERMOPLASTICO", "ISOLANTE")),
    ("Fundição e não ferrosos", ("FUNDICAO", "NAO FERROSO", "BRONZE", "ALUMINIO", "LATAO")),
    ("Veículos, pneus e auto peças", ("PNEU", "AUTO PECAS", "MATERIAL FERROVIARIO")),
    ("Cabos, correntes e içamento", ("CABO DE ACO", "LEVANTAMENTO", "EMPILHADEIRA", "FITA ACO")),
    ("Máquinas e equipamentos", ("MAQUINA", "EQUIPAMENTO", "PRENSA", "PERENSA", "EXAUSTOR", "VENTILADOR", "SOPRADOR", "CALANDRA", "LOCACAO")),
    ("Construção civil", ("CONSTRUCAO", "CIMENTO", "CONCRETO", "AREIA", "ARGAMASSA", "PEDRA", "VIDRO", "MADEIRA", "MARMORARIA")),
    ("Consumo, escritório e EPI", ("EMBALAGEM", "PALLET", "EPI", "UNIFORME", "LIMPEZA", "PAPELARIA", "ESCRITORIO", "INFORMATICA", "IMPRESSAO", "ETIQUETA", "FITA", "ADESIVO", "QUIMICO", "UTILIDADE", "HOTEL", "AGENCIA", "CARTORIO", "OTICA", "ESTANTE", "PRATILEIRA")),
]
# Código sem nome: pelo começo do código (padrão do cadastro).
PREFIXOS: list[tuple[str, str]] = [
    ("SV", "Serviços"), ("EF", "Fixação (parafusos, arruelas, anéis)"), ("SD", "Solda e gases"), ("TT", "Tintas"),
    ("TB", "Tubos"), ("TQ", "Tubos"), ("CH", "Chapas, bobinas e telas"), ("TO", "Chapas, bobinas e telas"), ("GRC", "Chapas, bobinas e telas"), ("BR", "Barras, perfis e vigas"), ("FJ", "Barras, perfis e vigas"), ("TMC", "Tubos"),
    ("SCOC", "Serviços"), ("SCT", "Serviços"),
    ("FRU", "Ferramentas e medição"), ("FRD", "Ferramentas e medição"), ("FRM", "Ferramentas e medição"),
    ("CLM", "Clientes"), ("FL", "Hidráulica, pneumática e conexões"), ("CNX", "Hidráulica, pneumática e conexões"),
]
OUTROS = "Outros"


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").upper()


def _texto(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v).replace("\x00", "")).strip()
    return s or None


def _nomes_codigos(linhas_fm) -> dict[str, str]:
    """{código: nome} da aba FM, abrindo "CLA*CLC*CLI=CHAPA ALUMINIO/CARBONO/INOX"."""
    nomes: dict[str, str] = {}
    for cod, nome in linhas_fm:
        cod, nome = _texto(cod), _texto(nome)
        if not cod or not nome:
            continue
        codigos = [c.strip().upper() for c in cod.split("*") if c.strip()]
        partes = [p.strip() for p in nome.split("/")]
        if len(codigos) > 1 and len(partes) == len(codigos):
            # "CHAPA ALUMINIO" / "CARBONO" / ... → base "CHAPA" + variante
            base = partes[0].rsplit(" ", 1)[0] if " " in partes[0] else ""
            variantes = [partes[0].rsplit(" ", 1)[-1], *partes[1:]]
            for c, v in zip(codigos, variantes):
                nomes.setdefault(c, f"{base} {v}".strip().upper())
        else:
            for c in codigos:
                nomes.setdefault(c, nome.upper())
    for c, n in NOMES_EXTRAS.items():
        nomes.setdefault(c.upper(), n)
    return nomes


def grupo_do_codigo(codigo: str, nome: str | None) -> str:
    if nome:
        alvo = _sem_acento(nome)
        for grupo, chaves in GRUPOS:
            if any(k in alvo for k in chaves):
                return grupo
    for prefixo, grupo in PREFIXOS:
        if codigo.startswith(prefixo):
            return grupo
    return OUTROS


def _fone(v) -> str | None:
    s = _texto(v)
    if not s or not re.search(r"\d{4}", s):
        return None
    return s


def _email(v) -> str | None:
    s = _texto(v)
    if not s or "@" not in s:
        return None
    return s.replace(" ", "")


def _cep(v) -> str | None:
    s = re.sub(r"\D", "", str(v or ""))
    if not s or not int(s):
        return None
    s = s.zfill(8)
    return f"{s[:5]}-{s[5:]}"


def ler_planilha(conteudo: bytes) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    if "FN" not in wb.sheetnames:
        raise ValueError(f"Aba 'FN' não encontrada. Abas: {', '.join(wb.sheetnames)}")
    nomes = _nomes_codigos(wb["FM"].iter_rows(min_row=2, max_col=2, values_only=True)) if "FM" in wb.sheetnames else {}
    fornecedores = []
    usados: dict[str, int] = {}
    for n, row in enumerate(wb["FN"].iter_rows(min_row=2, max_col=18, values_only=True), start=2):
        row = list(row) + [None] * (18 - len(row))
        nome = _texto(row[2])
        if not nome:
            continue
        codigos = [c.strip().upper() for c in str(row[1] or "").split("*") if c.strip()]
        for c in codigos:
            usados[c] = usados.get(c, 0) + 1
        contato = _texto(row[8])
        fornecedores.append({
            "l": n,  # linha na planilha
            "nome": nome,
            "cod": codigos,
            "fones": [f for f in (_fone(row[3]), _fone(row[4])) if f],
            "cel": [f for f in (_fone(row[5]), _fone(row[6])) if f],
            "email": _email(row[7]),
            "contato": None if not contato or contato.upper() in {"XX", "X", "-"} else contato,
            "marcas": _texto(row[9]),
            "materiais": _texto(row[10]),
            "endereco": _texto(row[13]),
            "cidade": _texto(row[14]),
            "cep": _cep(row[15]),
            "uf": (_texto(row[16]) or "").upper()[:2] or None,
            "cnpj": _texto(row[17]) if row[17] and re.search(r"\d{2}\.\d{3}", str(row[17])) else None,
        })
    wb.close()
    fornecedores.sort(key=lambda f: _sem_acento(f["nome"]))
    grupos = {c: {"nome": nomes.get(c), "grupo": grupo_do_codigo(c, nomes.get(c)), "n": q} for c, q in usados.items()}
    return {"fornecedores": fornecedores, "grupos": grupos}


def sincronizar(conteudo: bytes, caminho: str, modificado_em: datetime | None, forcar: bool = False) -> dict:
    """Grava se o arquivo mudou (hash)."""
    from psycopg.types.json import Jsonb

    sha = hashlib.sha256(conteudo).hexdigest()
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute("select arquivo_sha256 from fornecedores_planilha where id = 1")
        atual = cur.fetchone()
    if atual and atual[0] == sha and not forcar:
        return {"alterado": False}
    dados = ler_planilha(conteudo)
    with _conectar() as conn, conn.cursor() as cur:
        cur.execute(
            """insert into fornecedores_planilha (id, dados, arquivo, arquivo_sha256, arquivo_modificado_em, gerado_em)
               values (1, %s, %s, %s, %s, %s)
               on conflict (id) do update set dados = excluded.dados, arquivo = excluded.arquivo,
                 arquivo_sha256 = excluded.arquivo_sha256, arquivo_modificado_em = excluded.arquivo_modificado_em,
                 gerado_em = excluded.gerado_em""",
            (Jsonb(dados), caminho, sha, modificado_em, datetime.now(timezone.utc)),
        )
        conn.commit()
    return {"alterado": True, "fornecedores": len(dados["fornecedores"])}

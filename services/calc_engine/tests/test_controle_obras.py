import io
from datetime import datetime

import openpyxl
import pytest

from app.controle_obras import ler_planilha

CABECALHOS = [" PO", "CL   ", "ST", "QT", "Desenho", "RV", "Descrição", "Cor", "Cor3", "Tipar", "MAC", "META",
              "OBS", "PZ-C", "LARG", "Comp", "Alt", "Kg/pç", "Kg tot", "NF", "PPµ"]


def _planilha(linhas, aba="OBRAS") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    ws["A2"] = "E"  # colunas antes de H não entram
    for i, cab in enumerate(CABECALHOS):
        ws.cell(2, 8 + i, cab)
    for r, valores in enumerate(linhas, start=3):
        for i, v in enumerate(valores):
            ws.cell(r, 8 + i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_le_so_h_ate_ab_com_tipos():
    linha = ["4501760151-020", "W1", "A", 2.0, "A202998", "4", "CONE", "SP", None, "PV066719", "839.26",
             datetime(2026, 9, 23), "PINTURA", datetime(2026, 12, 16), 55, 176.0, 40.5, 8.5, 17.0, 7972, "SP"]
    colunas, linhas = ler_planilha(_planilha([linha, [None] * 21, linha]))
    assert [c["letra"] for c in colunas][:2] == ["H", "I"] and colunas[-1]["letra"] == "AB"
    assert colunas[0] == {"letra": "H", "cabecalho": "PO", "campo": "po", "tipo": "codigo"}
    assert colunas[-1]["campo"] == "ppu" and colunas[17]["campo"] == "kg_pc"
    assert [n for n, _ in linhas] == [3, 5]  # linha vazia ignorada, número da linha preservado
    v = linhas[0][1]
    assert v[3] == 2 and isinstance(v[3], int)  # 2.0 → 2
    assert v[11] == "2026-09-23" and v[13] == "2026-12-16"
    assert v[16] == 40.5 and v[19] == 7972


def test_aba_ausente_da_erro_claro():
    with pytest.raises(ValueError, match="OBRAS"):
        ler_planilha(_planilha([], aba="Outra"))

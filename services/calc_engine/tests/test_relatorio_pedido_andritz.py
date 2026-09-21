import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from app.relatorio_pedido_andritz import gerar_excel_pedido_andritz

ITEM = {
    "item": "4505093989-010",
    "valor_total": "2.866,68",
    "quantidade": 4,
    "material": "301947766",
    "material_antigo": "GPS-24A-TU-0613P319",
    "descricao": "SUPORTE DO SENSOR S355 J2",
    "mac": "792.26",
    "data_entrega": "24/08/26",
}


def test_gerar_excel_pedido_andritz_colunas_e_valores():
    import io

    conteudo = gerar_excel_pedido_andritz("4505093989", [ITEM])
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    ws = wb.active

    cabecalho = [c.value for c in ws[1]]
    assert cabecalho == ["Item", "Valor", "Quantidade", "Material", "Descrição", "MAC", "Data de entrega"]

    linha = [c.value for c in ws[2]]
    assert linha == [
        "4505093989-010",
        "2.866,68",
        4,
        "GPS-24A-TU-0613P319",
        "SUPORTE DO SENSOR S355 J2",
        "792.26",
        "24/08/26",
    ]

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from app.relatorio_pedido_weir import gerar_excel_pedido_weir

ITEM = {
    "item": "4501751360-010",
    "codigo": "A15792",
    "valor_total": 1050.01,
    "quantidade": 2,
    "descricao": "PROT.EC. DE CORRE",
    "referencia": "690.25.20.05",
    "data_entrega": "23/10/26",
}


def test_gerar_excel_pedido_weir_colunas_e_valores():
    import io

    conteudo = gerar_excel_pedido_weir([ITEM])
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    ws = wb.active

    cabecalho = [c.value for c in ws[1]]
    assert cabecalho == ["Item", "Código", "Valor", "Quantidade", "Descrição", "Referência", "Data de entrega"]

    linha = [c.value for c in ws[2]]
    assert linha == [
        "4501751360-010",
        "A15792",
        1050.01,
        2,
        "PROT.EC. DE CORRE",
        "690.25.20.05",
        "23/10/26",
    ]
    assert ws["C2"].number_format == '"R$" #,##0.00'

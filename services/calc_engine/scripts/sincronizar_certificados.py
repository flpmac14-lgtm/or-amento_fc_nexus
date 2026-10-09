"""Abas QUALIDADE e Backup Recebimento — pedido explícito do usuário: lista
dos certificados de matéria-prima do Recebimento, com NRI, descrição do
produto, tipo do material (tinta, insumo de solda etc.) e tipo de certificado
(ultrassom, LP, certificado de material…); e o sininho avisa quando um PDF é
adicionado, alterado ou excluído na pasta.

Duas pastas (FONTES): a principal ("1. Certificados de Matéria-Prima_Ordenados
pelo NRI", aba QUALIDADE) e o backup ("BACKUP RECEBIMENTO 20260828", aba
Backup Recebimento).

De onde vem cada coluna (estudado nos PDFs em 09/10/2026):
- NRI: nome do arquivo ("NRI 26-5097.pdf"; no ultrassom e no backup vem
  invertido, "NRI 0365-23 - CHAPA….pdf").
- Descrição e fornecedor: aba GERAL da planilha "CONTROLE DE RECEBIMENTO E
  ESTOQUE.xlsm" do Recebimento (uma linha por NRI); fora dela, o nome do arquivo.
- Tipo (material): código do material (coluna SECTRA = código do ERP,
  GGSSxxxx) → grupo/subgrupo do ERP (MT_SUBGRUPO, só SELECT). Sem código,
  pelas palavras da descrição.
- Tipo de certificado: conteúdo do PDF (texto ou OCR local), feito à parte por
  classificar_certificados.py (cache local); aqui só junta.

A pasta e o ERP só existem na rede da fábrica: roda nesta máquina (dentro do
ciclo de 15 min, sincronizar_controle_obras.py). A cada ciclo confere os
arquivos (nome, data, tamanho — é rápido); só relê a planilha e grava quando
algo mudou. Lista em qualidade_certificados (id 1 = principal, 2 = backup),
avisos em qualidade_notificacoes (migrations 0034/0036, sem policy). O site lê
por /api/qualidade/certificados?fonte=… (flpmac14 e perfil "qualidade").

Uso: python scripts/sincronizar_certificados.py [--simular] [--forcar]
"""

import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import warnings
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(RAIZ / ".env")

from app.orcamentos_salvos import _conectar  # noqa: E402
import classificar_certificados as classificar  # noqa: E402

RECEBIMENTO = Path(r"J:\5 - Almoxarifado - Recebimento\ISO_SGQ\Recebimento")
PLANILHA = RECEBIMENTO / "CONTROLE DE RECEBIMENTO E ESTOQUE.xlsm"

# id = linha em qualidade_certificados. subpastas: as que têm um PDF por NRI
# (na principal, as "NRI ANO 2014…2020" juntam dezenas de NRIs digitalizados
# num arquivo só — ficam fora). estado: última lista gravada (pros avisos), fora do git.
FONTES = {
    "principal": {
        "id": 1,
        "pasta": RECEBIMENTO / "1. Certificados de Matéria-Prima_Ordenados pelo NRI",
        "subpastas": ["", "ultrasson de chapas", "_REVISAR"],
        "estado": RAIZ / "sincronizar_certificados.estado.json",
    },
    "backup": {
        "id": 2,
        "pasta": RECEBIMENTO / "BACKUP RECEBIMENTO 20260828" / "NRI 20260828",
        "subpastas": [""],
        "estado": RAIZ / "sincronizar_certificados.backup.estado.json",
    },
}
# Compatibilidade (servir_certificados.py antigo).
PASTA = FONTES["principal"]["pasta"]
ESTADO = FONTES["principal"]["estado"]

# Mudança grande de uma vez (pasta renomeada, planilha reorganizada): um aviso
# só, em vez de milhares.
MAX_AVISOS = 150
ANO_ATUAL = datetime.now().year % 100

# Nome curto do tipo (pedido do usuário: "tinta, insumo de solda etc.") por
# grupo+subgrupo do código do ERP. O que não está aqui usa o nome do subgrupo.
TIPOS = {
    "5001": "Chapa",
    "5002": "Tubo",
    "5003": "Barra",
    "5004": "Perfil / viga",
    "5005": "Fundido / forjado",
    "5006": "Tela / manta",
    "5007": "Grade",
    "5008": "Tinta",
    "5009": "Retalho",
    "4004": "Insumo de solda",
    "4005": "Consumível de corte",
    "4006": "Consumível de jato / pintura",
    "4008": "Lubrificante / óleo",
    "4013": "Preparação de superfície",
    "4014": "Diluente / solvente",
    "4019": "Consumível de ensaio (LP / partículas)",
    "3001": "Elemento de fixação",
    "3002": "Transmissão / rolamento",
    "3004": "Vedação",
    "3005": "Revestimento",
    "3006": "Adesivo / fita",
    "3008": "Içamento",
}

# Sem código do ERP: palavras da descrição (a primeira que bater).
PALAVRAS = [
    (r"\bTINTA|PRIMER|ESMALTE|VERNIZ|CATALISADOR|\bCPB\b|\bCPA\b|SUMATANE|MACROPOXY|WEGTHANE|INTERGARD|INTERTHANE|REVRAN", "Tinta"),
    (r"DILUENTE|SOLVENTE|THINNER|REDUTOR", "Diluente / solvente"),
    (r"ARAME|ELETRODO|VARETA|FLUXO|\bE ?7018|\bER ?70|\bE ?71T|SOLDA", "Insumo de solda"),
    (r"PENETRANTE|REVELADOR|REMOVEDOR|MAGNAFLUX|MAGNAVIS|\bWCP|SPOTCHECK|METAL-?CHEK|PHENICON|PART[IÍ]CULA", "Consumível de ensaio (LP / partículas)"),
    (r"PARAFUSO|\bPRF\b|\bPRF\.|PORCA|ARRUELA|\bARR\b|PRISIONEIRO|CHUMBADOR|BARRA ROSCADA|PINO|CHAVETA|CUPILHA|REBITE|ANEL EL", "Elemento de fixação"),
    (r"FLANGE|CURVA|\bLUVA|NIPLE|REDU[CÇ][AÃ]O|\bTEE\b|BUJ[AÃ]O|\bPLUG|\bCOLAR\b|\bCAP\b|UNI[AÃ]O|COTOVELO|V[AÁ]LVULA", "Conexão / flange"),
    (r"\bCHAPA|\bCH\b|\bCH ?#|PLACA|\bINOX|\bFERRO\b|RET[AÂ]NGULO|\bDISCO|\bTAMPO|\bCFQ\b|LAMINADO", "Chapa"),
    (r"\bTUBO", "Tubo"),
    (r"\bBARRA|TARUGO|VERGALH", "Barra"),
    (r"CANTONEIRA|\bVIGA|PERFIL|\bU ?ENRIJ|\bTRILHO|\bPUDC|EXTRUDADO", "Perfil / viga"),
    (r"ROLAMENTO|MANCAL|ENGRENAGEM|ACOPLAMENTO", "Transmissão / rolamento"),
    (r"BUCHA|BRONZE|FUNDID|FORJAD", "Fundido / forjado"),
    (r"\bCABO\b|LINGA|MANILHA|OLHAL|SAPATILHO|TALHA|\bTROLE", "Içamento"),
    (r"\bTELA|LEN[CÇ]OL|FELTRO|\bMANTA", "Tela / manta"),
    (r"ADESIVO|\bFITA|\bCOLA\b|SCOTCHRAP", "Adesivo / fita"),
    (r"[OÓ]XIDO|GRANALHA|\bJATO", "Consumível de jato / pintura"),
    (r"JUNTA|VEDA[CÇ]|GAXETA|BORRACHA|O-?RING|VITON", "Vedação"),
]


def _nri_do_nome(nome: str) -> str | None:
    """'NRI 26-5097.pdf' → '26-5097'; 'NRI 0365-23 - CHAPA….pdf' → '23-0365'."""
    m = re.match(r"NRI\s*(\d{2})-(\d{3,5})", nome, re.I)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):04d}"
    m = re.search(r"NRI\D{0,4}(\d{3,4})-(\d{2})(?!\d)", nome, re.I)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):04d}"
    return None


def _descricao_do_nome(nome: str) -> str:
    """'NRI 1051-23 - ARAME MIG ER70S-6.pdf' → 'ARAME MIG ER70S-6'."""
    base = re.sub(r"\.pdf$", "", nome, flags=re.I)
    # Só tira o prefixo "NRI …" quando o nome começa com ele.
    m = re.match(r"(?:NRI\s*[\d-]+\s+)+-\s*(.+)$", base, re.I)
    texto = m.group(1) if m else base
    return re.sub(r"\s+", " ", re.sub(r"\s*-\s*C[oó]pia( \(\d+\))?$", "", texto, flags=re.I)).strip() or base


def _nri_da_planilha(v) -> str | None:
    m = re.search(r"(\d{2})-(\d{3,5})", str(v or ""))
    return f"{m.group(1)}-{int(m.group(2)):04d}" if m else None


def listar_pdfs(fonte: str) -> dict[str, dict]:
    """'sub/nome.pdf' → {nome, modificado, mtime, tamanho} da pasta da fonte."""
    cfg = FONTES[fonte]
    arquivos: dict[str, dict] = {}
    for sub in cfg["subpastas"]:
        base = cfg["pasta"] / sub if sub else cfg["pasta"]
        if not base.is_dir():
            continue
        with os.scandir(base) as it:
            for e in it:
                if e.is_file() and e.name.lower().endswith(".pdf"):
                    st = e.stat()
                    arquivos[f"{sub}/{e.name}" if sub else e.name] = {
                        "nome": e.name,
                        "modificado": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(timespec="seconds"),
                        "mtime": st.st_mtime,
                        "tamanho": st.st_size,
                    }
    return arquivos


def caminho_do_pdf(fonte: str, arquivo: str) -> Path:
    return FONTES[fonte]["pasta"] / arquivo


def todos_os_pdfs() -> list[tuple[Path, float, int]]:
    """(caminho, mtime, tamanho) de todos os PDFs das duas pastas — pro classificador."""
    saida = []
    for fonte, cfg in FONTES.items():
        if not cfg["pasta"].is_dir():
            continue
        for caminho, a in listar_pdfs(fonte).items():
            saida.append((caminho_do_pdf(fonte, caminho), a["mtime"], a["tamanho"]))
    return saida


def ler_planilha() -> dict[str, dict]:
    """NRI → {descricao, fornecedor, recebido, codigo} da aba GERAL."""
    import openpyxl

    # Cópia local: a planilha fica aberta no Recebimento (e na rede é lenta).
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "controle.xlsm"
        shutil.copyfile(PLANILHA, copia)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
            ws = wb["GERAL"]
            cab = [str(c or "").strip().upper() for c in next(ws.iter_rows(max_row=1, values_only=True))]
            col = {nome: cab.index(nome) for nome in ("NRI", "SECTRA", "DESC SECTRA", "DESCRICAO", "FORNECEDOR", "DT RECEB.")}
            linhas = list(ws.iter_rows(min_row=2, values_only=True))
            wb.close()

    def valor(r, nome):
        return r[col[nome]] if col[nome] < len(r) else None

    def txt(r, nome):
        v = valor(r, nome)
        return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else None

    por_nri: dict[str, dict] = {}
    for r in linhas:
        nri = _nri_da_planilha(valor(r, "NRI"))
        if not nri:
            continue
        desc = txt(r, "DESCRICAO") or txt(r, "DESC SECTRA")
        codigo = txt(r, "SECTRA")
        codigo = codigo if codigo and re.fullmatch(r"\d{8}", codigo) else None
        receb = valor(r, "DT RECEB.")
        item = por_nri.setdefault(nri, {"descricao": [], "fornecedor": None, "recebido": None, "codigo": None})
        if desc and desc not in item["descricao"]:
            item["descricao"].append(desc)
        item["fornecedor"] = item["fornecedor"] or txt(r, "FORNECEDOR")
        item["codigo"] = item["codigo"] or codigo
        if isinstance(receb, datetime) and not item["recebido"]:
            item["recebido"] = receb.date().isoformat()
    return por_nri


def subgrupos_erp() -> dict[str, str]:
    """'5008' → 'TINTA, VERNIZ, …' (MT_SUBGRUPO do ERP, só SELECT)."""
    import pyodbc

    conn = pyodbc.connect(
        "DRIVER={ODBC Driver 17 for SQL Server};"
        f"SERVER={os.environ['ERP_SQL_SERVER']};DATABASE={os.environ['ERP_SQL_DATABASE']};"
        f"UID={os.environ['ERP_SQL_USER']};PWD={os.environ['ERP_SQL_PASSWORD']};TrustServerCertificate=yes",
        timeout=10,
    )
    try:
        cur = conn.cursor()
        cur.execute("SELECT GRUPO, CODIGO, DESCRICAO FROM MT_SUBGRUPO")
        return {f"{g.strip()}{c.strip()}": (d or "").strip() for g, c, d in cur.fetchall()}
    finally:
        conn.close()


def tipo_de(codigo: str | None, descricao: str, subgrupos: dict[str, str]) -> str | None:
    if codigo:
        chave = codigo[:4]
        if chave in TIPOS:
            return TIPOS[chave]
        if subgrupos.get(chave):
            return subgrupos[chave].capitalize()
    texto = descricao.upper()
    for padrao, tipo in PALAVRAS:
        if re.search(padrao, texto):
            return tipo
    return None


class _Fontes:
    """Planilha e ERP lidos uma vez só por rodada (e só se alguma pasta mudou)."""

    def __init__(self):
        self._planilha = None
        self._subgrupos = None

    @property
    def planilha(self) -> dict:
        if self._planilha is None:
            self._planilha = ler_planilha()
        return self._planilha

    @property
    def subgrupos(self) -> dict:
        if self._subgrupos is None:
            try:
                self._subgrupos = subgrupos_erp()
            except Exception as e:  # noqa: BLE001 — sem ERP, usa só os nomes fixos e as palavras
                print(f"ERP indisponível ({e}); tipo só pelos nomes fixos/palavras.")
                self._subgrupos = {}
        return self._subgrupos


def montar(fonte: str, arquivos: dict[str, dict], dados: _Fontes, cache: dict) -> list[dict]:
    itens = []
    for caminho, arq in arquivos.items():
        nome = arq["nome"]
        nri = _nri_do_nome(nome)
        info = dados.planilha.get(nri) if nri else None
        descricao = " / ".join(info["descricao"]) if info and info["descricao"] else _descricao_do_nome(nome)
        certificado = classificar.tipos_do_cache(cache, caminho_do_pdf(fonte, caminho), arq["mtime"], arq["tamanho"])
        itens.append({
            "nri": nri,
            "descricao": descricao,
            "tipo": tipo_de(info["codigo"] if info else None, descricao, dados.subgrupos),
            "certificado": certificado,  # None = ainda não analisado
            "fornecedor": info["fornecedor"] if info else None,
            "recebido": info["recebido"] if info else None,
            "codigo": info["codigo"] if info else None,
            "arquivo": caminho,
            "modificado": arq["modificado"],
            "tamanho": arq["tamanho"],
            "na_planilha": bool(info),
        })

    def ordem(i):
        if not i["nri"]:
            return (1, 0, 0, i["arquivo"])
        ano, num = i["nri"].split("-")
        # Ano depois do atual (ex.: "NRI 0092-91", digitado errado) vai pro fim.
        return (0 if int(ano) <= ANO_ATUAL else 1, -int(ano), -int(num), i["arquivo"])

    itens.sort(key=ordem)  # mais novo primeiro; sem NRI no fim
    return itens


def _assinatura(arquivos: dict[str, dict]) -> str:
    """Muda quando a planilha, algum PDF (nome, data, tamanho) ou a classificação mudou."""
    base = json.dumps(sorted((k, v["modificado"], v["tamanho"]) for k, v in arquivos.items()))
    cache = classificar.CACHE.stat().st_mtime_ns if classificar.CACHE.exists() else 0
    return hashlib.sha256(f"{PLANILHA.stat().st_mtime_ns}|{cache}|{base}".encode()).hexdigest()


def _avisos(antes: list[dict], depois: list[dict]) -> list[tuple]:
    """(tipo, nri, arquivo, descricao) do que entrou, mudou ou saiu da pasta.
    O tipo de certificado não conta (ele chega aos poucos, pela classificação)."""
    a = {i["arquivo"]: i for i in antes}
    d = {i["arquivo"]: i for i in depois}
    avisos = []
    for k, i in d.items():
        if k not in a:
            avisos.append(("adicionado", i["nri"], k, i["descricao"]))
        elif (a[k].get("modificado"), a[k].get("tamanho"), a[k]["descricao"], a[k]["tipo"]) != (
            i["modificado"], i["tamanho"], i["descricao"], i["tipo"],
        ):
            avisos.append(("alterado", i["nri"], k, i["descricao"]))
    for k, i in a.items():
        if k not in d:
            avisos.append(("excluido", i["nri"], k, i["descricao"]))
    if len(avisos) > MAX_AVISOS:
        nomes = {"adicionado": "adicionados", "alterado": "alterados", "excluido": "excluídos"}
        partes = [f"{sum(1 for x in avisos if x[0] == t)} {n}" for t, n in nomes.items() if any(x[0] == t for x in avisos)]
        return [("varios", None, None, f"Mudança grande na pasta: {', '.join(partes)}.")]
    return avisos


def _ler_estado(caminho: Path) -> dict:
    try:
        return json.loads(caminho.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def sincronizar(simular: bool = False, forcar: bool = False) -> str | None:
    """Grava lista e avisos de cada pasta que mudou (None = nada mudou)."""
    if not PLANILHA.exists():
        return None  # fora da rede da fábrica / J: desconectado
    dados = _Fontes()
    cache = classificar.ler_cache()
    resumos = []
    for fonte, cfg in FONTES.items():
        if not cfg["pasta"].is_dir():
            continue
        arquivos = listar_pdfs(fonte)
        assinatura = _assinatura(arquivos)
        estado = _ler_estado(cfg["estado"])
        if not forcar and not simular and estado.get("assinatura") == assinatura:
            continue
        itens = montar(fonte, arquivos, dados, cache)
        sem_planilha = sum(1 for i in itens if not i["na_planilha"])
        analisados = sum(1 for i in itens if i["certificado"])
        # 1ª vez (sem lista anterior): não avisa — senão seriam 7 mil "adicionados".
        avisos = _avisos(estado["itens"], itens) if "itens" in estado else []
        resumo = (f"{fonte}: {len(itens)} certificados ({sem_planilha} fora da planilha, "
                  f"{analisados} com tipo de certificado), {len(avisos)} aviso(s)")
        if simular:
            resumos.append(resumo + " — simulado")
            continue
        agora = datetime.now(timezone.utc)
        with _conectar() as conn, conn.cursor() as cur:
            cur.execute(
                """insert into qualidade_certificados (id, dados, gerado_em) values (%s, %s, %s)
                   on conflict (id) do update set dados = excluded.dados, gerado_em = excluded.gerado_em""",
                (cfg["id"], json.dumps({"itens": itens}, ensure_ascii=False), agora),
            )
            if avisos:
                cur.executemany(
                    """insert into qualidade_notificacoes (fonte, tipo, nri, arquivo, descricao, criado_em)
                       values (%s, %s, %s, %s, %s, %s)""",
                    [(fonte, *a, agora) for a in avisos],
                )
        cfg["estado"].write_text(json.dumps({"assinatura": assinatura, "itens": itens}, ensure_ascii=False), encoding="utf-8")
        resumos.append(resumo + " — gravado")
    return " | ".join(resumos) or None


# Nome usado no ciclo de 15 min (sincronizar_controle_obras.py).
sincronizar_se_mudou = sincronizar


if __name__ == "__main__":
    print(sincronizar(simular="--simular" in sys.argv, forcar="--forcar" in sys.argv) or "sem mudança")
